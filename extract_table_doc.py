from docx import Document

p = r"C:\Users\user\iCloudDrive\5학기\연구미팅\VD PPT\20260822 Table.docx"
doc = Document(p)
for i, para in enumerate(doc.paragraphs):
    t = para.text.strip()
    if t:
        print(f"P{i}: {t}")
for ti, table in enumerate(doc.tables):
    print(f"\n===== TABLE {ti+1} {len(table.rows)}x{len(table.columns)} =====")
    for row in table.rows:
        print("\t".join(cell.text.replace("\n", " | ") for cell in row.cells))
