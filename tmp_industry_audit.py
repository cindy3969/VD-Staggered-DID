from pathlib import Path
from zipfile import ZipFile
from lxml import etree
import re, sys
import openpyxl
sys.stdout.reconfigure(encoding='utf-8')
base=Path('C:/Users/user/iCloudDrive/5학기/연구미팅')
if sys.argv[1]=='docs':
 for p in sorted((base/'VD PPT').glob('*.docx')):
  with ZipFile(p) as z:
   root=etree.fromstring(z.read('word/document.xml'))
   ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
   texts=[''.join(x.itertext()) for x in []]
   texts=[''.join(x.xpath('.//w:t/text()',namespaces=ns)) for x in root.xpath('//w:p',namespaces=ns)]
   hits=[t for t in texts if re.search(r'industry\s*[12]|industries|classif|software|deeptech|업종|산업',t,re.I)]
   hits=[t for t in hits if not re.match(r'^(reg |asdoc|psmatch|reghdfe)',t)]
   if hits: print('\nFILE',p.name,'\n'+'\n'.join(hits))
else:
 for name in sys.argv[2:]:
  p=base/'20240108-2010 to 2015 firms'/name
  w=openpyxl.load_workbook(p,read_only=True,data_only=False)
  print('FILE',name,'SHEETS',w.sheetnames,flush=True)
  for s in w:
   rows=s.iter_rows(values_only=True)
   head=next(rows,[])
   for attempt in range(8):
    if any(v and re.search('industr',str(v),re.I) for v in head): break
    head=next(rows,[])
   inds=[i for i,v in enumerate(head) if v and re.search('industr',str(v),re.I)]
   print('SHEET',s.title,s.max_row,'HEAD',[(i+1,head[i]) for i in inds],flush=True)
   if inds:
    from collections import Counter
    data=list(rows)
    c=Counter(tuple(str(row[i]) for i in inds) for row in data)
    print('UNIQUE',len(c),flush=True)
    for k,n in c.most_common(8): print(n,k)
    sub=[i for i in inds if re.search(r'industry\s*[12]',str(head[i]),re.I)]
    print('MAPPING',Counter(tuple(str(row[i]) for i in sub) for row in data))
