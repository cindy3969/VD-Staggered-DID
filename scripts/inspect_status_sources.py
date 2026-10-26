import re
from collections import Counter

import pandas as pd
from openpyxl import load_workbook

dta_path = r"C:\Users\user\Documents\VD\outputs\psm_final_2820_package\table3_current_2820_cross_section_v2.dta"
xlsx_path = r"C:\Users\user\Documents\VD\2010 2015 Crunchbase deals merged_v4_series-updated.xlsx"

dta = pd.read_stata(dta_path, convert_categoricals=False)
norm = lambda value: re.sub(r"[^a-z0-9]", "", str(value).lower())

wb = load_workbook(xlsx_path, read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
rows = []
for row in ws.iter_rows(min_row=3, values_only=True):
    url, name, status = row[1], row[3], row[12]
    if name:
        rows.append((url, name, status))

xlsx_names = {norm(name) for _, name, _ in rows}
dta_names = set(dta["OrganizationName"].map(norm))
print("source rows", len(rows), flush=True)
print("source unique", len(xlsx_names), flush=True)
print("status counts", Counter(status for _, _, status in rows), flush=True)
print("exact normalized coverage", sum(name in xlsx_names for name in dta_names), "of", len(dta_names), flush=True)
print("dta unique", dta["OrganizationName"].nunique(), flush=True)
print("urls unique", len({url for url, _, _ in rows if url}), flush=True)
print(
    "unmatched sample",
    dta.loc[~dta["OrganizationName"].map(norm).isin(xlsx_names), "OrganizationName"].head(50).tolist(),
    flush=True,
)
