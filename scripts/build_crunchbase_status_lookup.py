import re
from collections import defaultdict

import pandas as pd
from openpyxl import load_workbook

ROOT = r"C:\Users\user\Documents\VD"
PACKAGE = ROOT + r"\outputs\psm_final_2820_package"
SOURCE_XLSX = ROOT + r"\2010 2015 Crunchbase deals merged_v4_series-updated.xlsx"
TARGET_DTA = PACKAGE + r"\table3_current_2820_cross_section_v2.dta"
LEGACY_DTA = PACKAGE + r"\ORIGINAL_final_matched_2820.dta"
OUT_DTA = PACKAGE + r"\_crunchbase_status_lookup.dta"
OUT_CSV = PACKAGE + r"\crunchbase_status_update_audit.csv"


def key(value):
    return re.sub(r"[^a-z0-9]", "", str(value).casefold())


target = pd.read_stata(TARGET_DTA, convert_categoricals=False)
legacy = pd.read_stata(LEGACY_DTA, convert_categoricals=False)
legacy["name_key"] = legacy["OrganizationName"].map(key)
legacy = legacy.drop_duplicates("name_key").set_index("name_key")

wb = load_workbook(SOURCE_XLSX, read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
headers = [cell.value for cell in next(ws.iter_rows(min_row=2, max_row=2))]
idx = {name: i for i, name in enumerate(headers) if name is not None}

needed = [
    "Organization Name URL",
    "Organization Name",
    "Funding Status",
    "Operating Status",
    "IPO Status",
    "IPO Date",
]
missing = [name for name in needed if name not in idx]
if missing:
    raise RuntimeError(f"Missing Crunchbase columns: {missing}")

source = defaultdict(dict)
for row in ws.iter_rows(min_row=3, values_only=True):
    name = row[idx["Organization Name"]]
    if not name:
        continue
    k = key(name)
    rec = source[k]
    for field in needed:
        value = row[idx[field]]
        if value not in (None, "") and field not in rec:
            rec[field] = value

records = []
for _, row in target.iterrows():
    k = key(row["OrganizationName"])
    rec = source.get(k)
    if rec is None:
        raise RuntimeError(f"No Crunchbase source match: {row['OrganizationName']}")

    funding_status = str(rec.get("Funding Status", "")).strip()
    operating_status = str(rec.get("Operating Status", "")).strip()
    ipo_status = str(rec.get("IPO Status", "")).strip()
    ipo_date = pd.to_datetime(rec.get("IPO Date"), errors="coerce")
    ipo_year = float(ipo_date.year) if not pd.isna(ipo_date) else None
    exact_event_date = ipo_date.strftime("%Y-%m-%d") if not pd.isna(ipo_date) else ""

    legacy_type = ""
    legacy_year = None
    if k in legacy.index:
        legacy_type = str(legacy.at[k, "exittype2"]).strip().lower()
        value = legacy.at[k, "exityear2"]
        legacy_year = None if pd.isna(value) else float(value)

    vd_year = float(row["VDyear"])
    timing_source = "current Crunchbase snapshot"
    event_year = None

    if funding_status.upper() == "M&A":
        status = "merged" if legacy_type == "merged" else "acquired"
        exact_event_date = ""
        if legacy_type in {"acquired", "merged"} and legacy_year is not None:
            event_year = legacy_year
            timing_source = "Crunchbase event year"
    elif funding_status.upper() == "IPO" or ipo_status.lower() == "public":
        status = "public"
        event_year = ipo_year
        timing_source = "Crunchbase IPO date" if ipo_year is not None else timing_source
    elif operating_status.lower() == "closed":
        status = "defunct"
        exact_event_date = ""
        if legacy_type == "closed" and legacy_year is not None:
            event_year = legacy_year
            timing_source = "Crunchbase event year"
    else:
        status = "private"
        exact_event_date = ""

    # When Crunchbase supplies an event year, only post-reference-year exits count.
    # Static current statuses without a dated event remain classified from the current snapshot.
    if event_year is not None and event_year <= vd_year:
        status = "private"
        event_year = None
        exact_event_date = ""

    records.append(
        {
            "firm_id": int(row["firm_id"]),
            "status": status,
            "m_a": int(status in {"acquired", "merged"}),
            "ipo": int(status == "public"),
            "defunct": int(status == "defunct"),
            "status_year": event_year,
            "status_date": exact_event_date,
            "status_event_year": event_year,
            "status_timing_source": timing_source,
            "crunchbase_url": str(rec.get("Organization Name URL", "")),
            "crunchbase_funding_status": funding_status,
            "crunchbase_operating_status": operating_status,
            "crunchbase_ipo_status": ipo_status,
        }
    )

out = pd.DataFrame(records)
if len(out) != 2820 or out["firm_id"].nunique() != 2820:
    raise RuntimeError("Status lookup is not a unique 2,820-firm file")
if out["status"].isna().any():
    raise RuntimeError("Missing status values")

out[["firm_id", "status", "m_a", "ipo", "defunct", "status_year", "status_date"]].to_stata(
    OUT_DTA,
    write_index=False,
    version=118,
    variable_labels={
        "status": "Post-VD/current firm status from Crunchbase",
        "m_a": "1 if acquired or merged; 0 otherwise",
        "ipo": "1 if public; 0 otherwise",
        "defunct": "1 if defunct/closed; 0 otherwise",
        "status_year": "Post-VD exit/event year from Crunchbase",
        "status_date": "Exact Crunchbase event date (YYYY-MM-DD), when available",
    },
)
out.to_csv(OUT_CSV, index=False, encoding="utf-8-sig")

print(out["status"].value_counts().to_string())
print("dummy sums", out[["m_a", "ipo", "defunct"]].sum().to_dict())
print("dated events", int(out["status_event_year"].notna().sum()))
print("rows", len(out), "unique firm_id", out["firm_id"].nunique())
