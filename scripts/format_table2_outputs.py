from pathlib import Path
import json
import math

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(r"C:\Users\user\Documents\VD")
OUT = ROOT / "outputs" / "table2_forward_citations_application_level"

rows = [pd.read_csv(OUT / f"table2_boot_{i}.csv").iloc[0] for i in range(5)]
res = pd.DataFrame(rows)

def pvalue(est, se):
    return math.erfc(abs(float(est) / float(se)) / math.sqrt(2)) if se and not math.isnan(se) else math.nan

for stem, se_stem, p_stem in [("simple", "simple_se", "simple_p"), ("post_avg", "post_avg_se", "post_avg_p")]:
    res[p_stem] = [pvalue(a, b) for a, b in zip(res[stem], res[se_stem])]
for e in range(-4, 5):
    res[f"p{e}"] = [pvalue(a, b) for a, b in zip(res[f"e{e}"], res[f"se{e}"])]
res.to_csv(OUT / "Table_2_adjusted_forward_citation_results.csv", index=False, encoding="utf-8-sig")

labels = ["All firms", "No prior patent", "Prior patent", "Early stage", "Late stage"]

def stars(p):
    return "***" if p < .01 else "**" if p < .05 else "*" if p < .10 else ""

def fmt(est, p):
    return f"{est:.3f}{stars(p)}"

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Inches(11.69), Inches(8.27)
sec.top_margin = sec.bottom_margin = Inches(0.65)
sec.left_margin = sec.right_margin = Inches(0.65)

styles = doc.styles
styles["Normal"].font.name = "Times New Roman"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
styles["Normal"].font.size = Pt(9)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Table 2. Staggered DiD Estimates on Forward Citations")
r.bold = True
r.font.size = Pt(13)
r.font.color.rgb = RGBColor(0, 0, 0)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Dependent variable: Mean filing-year-adjusted forward citations per granted patent application")
r.italic = True
r.font.size = Pt(9)

table = doc.add_table(rows=1, cols=6)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [2.15, 1.25, 1.25, 1.25, 1.25, 1.25]
for i, w in enumerate(widths):
    table.columns[i].width = Inches(w)
hdr = table.rows[0].cells
hdr[0].text = ""
for i, lab in enumerate(labels, 1):
    hdr[i].text = f"({i})\n{lab}"

def add_line(label, values, italic=False):
    cells = table.add_row().cells
    cells[0].text = label
    for i, v in enumerate(values, 1):
        cells[i].text = str(v)
    for c in cells:
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for pp in c.paragraphs:
            pp.alignment = WD_ALIGN_PARAGRAPH.CENTER if c is not cells[0] else WD_ALIGN_PARAGRAPH.LEFT
            for rr in pp.runs:
                rr.italic = italic
                rr.font.size = Pt(8.5)

add_line("Simple ATT", [fmt(r.simple, r.simple_p) for _, r in res.iterrows()])
add_line("", [f"({r.simple_se:.3f})" for _, r in res.iterrows()], italic=True)
add_line("Aggregated event-study ATT", [fmt(r.post_avg, r.post_avg_p) for _, r in res.iterrows()])
add_line("", [f"({r.post_avg_se:.3f})" for _, r in res.iterrows()], italic=True)
add_line("Controls", ["Yes"] * 5)
add_line("Region FE", ["Yes"] * 5)
add_line("Industry FE", ["Yes"] * 5)
add_line("Comparison group", ["Never treated"] * 5)
add_line("Observations", [f"{int(r.obs):,}" for _, r in res.iterrows()])
add_line("Firms", [f"{int(r.firms):,}" for _, r in res.iterrows()])

for row_i, row in enumerate(table.rows):
    for cell in row.cells:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "D9EAF7" if row_i == 0 else ("F5F8FA" if row_i % 2 == 0 else "FFFFFF"))
        tcPr.append(shd)
        tcMar = tcPr.first_child_found_in("w:tcMar")
        if tcMar is None:
            tcMar = OxmlElement("w:tcMar"); tcPr.append(tcMar)
        for side in ["top", "left", "bottom", "right"]:
            node = OxmlElement(f"w:{side}"); node.set(qn("w:w"), "90"); node.set(qn("w:type"), "dxa"); tcMar.append(node)
        for pp in cell.paragraphs:
            if row_i == 0:
                pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for rr in pp.runs: rr.bold = True

note = doc.add_paragraph()
note.paragraph_format.space_before = Pt(6)
note.paragraph_format.line_spacing = 1.0
note.add_run("Note: ").bold = True
note.add_run(
    "Forward citations are counted through 2024 at the granted patent-application level. "
    "Following the supplied paper's truncation-adjustment logic, raw citations are divided by the mean citation count of applications filed in the same year. "
    "CPC/technology-class cells are intentionally not used. Firm-year outcomes are the mean adjusted citations of ultimately granted applications filed in that year, with zero in firm-years without such applications. "
    "Estimation uses Callaway-Sant'Anna doubly robust staggered DiD, never-treated controls, the Table 1 controls and fixed effects, and firm bootstrap standard errors (49 replications). "
    "The aggregated event-study ATT averages event times 0 through 4. Standard errors are in parentheses. *, **, and *** denote significance at the 10%, 5%, and 1% levels."
)

lim = doc.add_paragraph()
lim.paragraph_format.space_before = Pt(4)
lim.add_run("Data note: ").bold = True
lim.add_run(
    "Of 6,319 grant-linked applications, 5,801 retain a recovered citation value from the prior 2024 snapshot. "
    "The remaining 518 have no recovered citation-link row and are coded as zero under the citation-table left-join convention. "
    "The accompanying application-level file preserves an indicator for this distinction."
)

doc.save(OUT / "Table_2_forward_citations_replication.docx")

readme = f"""TABLE 2 FORWARD-CITATION REPLICATION\n\nPrimary method\n- Unit: granted patent application.\n- Citation cutoff: 2024.\n- Right-truncation adjustment: raw forward citations / mean raw citations among linked granted applications in the same filing year.\n- CPC/technology class: not used, following the project decision.\n- Firm-year DV: mean adjusted citations across ultimately granted applications filed in the firm-year; zero when none.\n- Estimator: Callaway-Sant'Anna doubly robust staggered DiD, never-treated controls, Table 1 controls/FE, 49 firm bootstrap replications.\n\nCoverage\n- Applications: 10,727\n- Grant-linked applications: 6,319\n- Recovered citation values: 5,801\n- No recovered citation-link row, coded zero: 518\n\nFiles\n- application_level_forward_citations_through_2024.csv: analysis-ready application-level citation data.\n- citation_normalization_by_filing_year.csv: annual denominators.\n- firm_year_panel_with_forward_citations.csv: estimation panel.\n- Table_2_adjusted_forward_citation_results.csv: full point estimates, bootstrap SEs and p-values.\n- table2_raw_sensitivity.csv: raw-citation point-estimate sensitivity analysis.\n- Table_2_forward_citations_replication.docx: formatted table.\n- dripw_table2.py: estimator.\n- QA_summary.json: reproducibility checks.\n"""
(OUT / "README.txt").write_text(readme, encoding="utf-8")

print(res[["sub", "simple", "simple_se", "simple_p", "post_avg", "post_avg_se", "post_avg_p", "obs", "firms"]].to_string(index=False))
