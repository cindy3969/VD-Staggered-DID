from pathlib import Path
from copy import deepcopy
from io import BytesIO
import zipfile,json,hashlib,re
import pandas as pd
from docx import Document
from docx.table import Table
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches,Pt

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/table_run_v2_20261020'
SOURCE=ROOT/'outputs/table_update_20261019/VD Table 20261019.docx'
FINAL=OUT/'VD Table 20261020 original 2820 v2.docx'
D=Document(SOURCE); templates=[deepcopy(t._tbl) for t in D.tables]
C=pd.read_stata(OUT/'coefficients.dta')
assert C.model.nunique()==97
assert 'WORD_TABLES_RUN_COMPLETE' in (OUT/'RUN_TABLE3_FULL_v2.log').read_text(errors='replace')
META={m['model']:m for m in json.loads((OUT/'models.json').read_text())}
body=D._element.body
caption=deepcopy(next(p._p for p in D.paragraphs if p.text.startswith('TABLE 3.')))
note=deepcopy(next(p._p for p in D.paragraphs if p.text.startswith('Note: The DV is ln(1 + cohort')))
portrait=deepcopy(D.sections[0]._sectPr)
landscape=deepcopy(D.sections[1]._sectPr)
for child in list(body):body.remove(child)
body.append(deepcopy(portrait))
audit=[]; inventory=[]; current_landscape=False; count=0

def paragraph(el,text):
    texts=el.findall('.//'+qn('w:t'))
    if texts:
        texts[0].text=text
        for t in texts[1:]:t.text=''
    else:
        r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=text;r.append(t);el.append(r)
    # Do not inherit obsolete section boundaries or page-break runs.
    for x in list(el.findall('.//'+qn('w:sectPr')))+list(el.findall('.//'+qn('w:br'))):
        x.getparent().remove(x)
    return el

def cell(t,r,c,text):
    target=t.cell(r,c)
    paragraph(target.paragraphs[0]._p,str(text))
    for p in target.paragraphs[1:]:paragraph(p._p,'')

def num(x,integer=False):
    if pd.isna(x):return '—'
    if integer:return f'{int(x):,}'
    return f'{0 if abs(x)<.0005 else x:.3f}'

def entry(mid,term):
    z=C[C.model.eq(mid)]
    z=z[z.term.str.split(':').str[-1].isin([term,'1.'+term]) & z.se.gt(0)]
    assert len(z)<=1,(mid,term,z)
    return None if z.empty else z.iloc[0]

def put(t,r,col,mid,term):
    z=entry(mid,term)
    b=p=''
    if z is not None:
        b=num(z.b)+('***' if z.p<.01 else '**' if z.p<.05 else '*' if z.p<.1 else '')
        p='(<0.001)' if z.p<.001 else '('+num(z.p)+')'
    cell(t,r,col,b);cell(t,r+1,col,p)
    audit.append(dict(model=mid,term=term,b=b,p=p))

def add(capt,t,notes,wide=False):
    global count,current_landscape
    if count:
        if wide!=current_landscape:
            p=OxmlElement('w:p');pr=OxmlElement('w:pPr')
            s=deepcopy(landscape if current_landscape else portrait)
            typ=s.find(qn('w:type'))
            if typ is None:typ=OxmlElement('w:type');s.append(typ)
            typ.set(qn('w:val'),'nextPage');pr.append(s);p.append(pr)
            body.insert(len(body)-1,p)
        else:
            p=OxmlElement('w:p');r=OxmlElement('w:r');br=OxmlElement('w:br')
            br.set(qn('w:type'),'page');r.append(br);p.append(r);body.insert(len(body)-1,p)
    current_landscape=wide
    body.remove(body[-1]);body.append(deepcopy(landscape if wide else portrait))
    cap=paragraph(deepcopy(caption),capt)
    body.insert(len(body)-1,cap)
    # Remove fixed heights and prevent individual data rows splitting.
    for row in t.rows:
        for h in row._tr.findall('.//'+qn('w:trHeight')):h.getparent().remove(h)
        pr=row._tr.get_or_add_trPr()
        if pr.find(qn('w:cantSplit')) is None:pr.append(OxmlElement('w:cantSplit'))
    head=t.rows[0]._tr.get_or_add_trPr()
    if head.find(qn('w:tblHeader')) is None:head.append(OxmlElement('w:tblHeader'))
    body.insert(len(body)-1,t._tbl)
    n=paragraph(deepcopy(note),notes)
    body.insert(len(body)-1,n)
    inventory.append(capt);count+=1

def clone(i):return Table(deepcopy(templates[i]),D._body)
common=('Prior Patent = 1 for firms with prior patents and 0 otherwise (priorpatentdummy_VD). Controls are measured at the VD/reference year. '
        'Two-sided p-values are in parentheses. * p<0.10, ** p<0.05, *** p<0.01, based on unrounded p-values.')
terms={2:'VD',4:'1.VD#1.priorpatentdummy_VD',6:'1.VD#1.Earlystage_VD',
       8:'priorpatentdummy_VD',10:'Earlystage_VD',12:'ln_MoneyRaised_VD',14:'NumofFundingRounds_VD',16:'_cons'}
labels={2:'VD',4:'VD × Prior Patent',6:'VD × Early Stage',8:'Prior Patent',
        10:'Early Stage',12:'Log(Total Funding at VD)',14:'Funding Rounds at VD',16:'Constant'}

def regtable(capt,models,dv,notes='',template=3,cox=False):
    t=clone(template);cell(t,0,0,dv)
    for r,term in terms.items():
        if cox and r==16:continue
        cell(t,r,0,labels[r])
        for col,mid in enumerate(models,1):put(t,r,col,mid,term)
    for col,mid in enumerate(models,1):
        z=C[C.model.eq(mid)].iloc[0]
        cell(t,21 if cox else 22,col,num(z.N,True))
        cell(t,20 if cox else 23,col,num(z.fail,True) if cox else num(z.r2))
    add(capt,t,notes+' '+common)
    return t

# Descriptive and correlation tables follow the exact active variable list.
S=pd.read_stata(OUT/'summary.dta').set_index('variable')
R=pd.read_stata(OUT/'correlations.dta').set_index(['variable1','variable2'])
vs=list(S.index)
ls=['Patent Applications (3-year filing window)','VD','Prior Patent','Early Stage',
    'Total Funding at VD (USD millions)','Funding Rounds at VD','Bio/Healthcare',
    'Hardware/Manufacturing/Deep Tech','Software/IT services','California','Massachusetts','New York']
t=clone(0)
for r,v in enumerate(vs,1):
    cell(t,r,0,ls[r-1])
    values=S.loc[v,['N','mean','sd','min','max']].to_numpy(copy=True)
    if v=='MoneyRaised_VD':values[1:]/=1e6
    for c,val in enumerate(values,1):cell(t,r,c,num(val,c==1))
add('Table 1. Summary Statistics',t,
    'Original 2,820-firm sample from 260122 winsorization.dta, not the rematched sample. '
    'The summary DV is patent_application_win3, as specified in the do-file; it differs from the normalized log DV in Table 3. '
    'Funding is in USD millions. Indicators are proportions.')
for idx,rows,offset in [(1,12,0),(2,6,6)]:
    t=clone(idx)
    for r in range(1,rows+1):
        i=r-1+offset;cell(t,r,1,ls[i])
        for col in range(2,8):
            j=col-2+(0 if idx==1 else 6)
            z=R.loc[(vs[i],vs[j])]
            value='' if i<j else '1.000' if i==j else num(z.r)+('*' if z.p<.05 else '')
            cell(t,r,col,value)
    add(f'Table 2. Pairwise Correlations — Panel {"A" if idx==1 else "B"}',t,
        '* p<0.05. Pairwise correlations use the Table 1 variables and original 2,820-firm sample.',wide=True)
regtable('Table 3. VD Backing and Post-VD Patent Quantity',range(1,5),'DV: Log(1 + Normalized Granted Applications)',
         'DV: ln(1 + grantedpatentbyfiling_norm_win3), filing years t+1 through t+3. OLS with absorbed region, industry and founding-year fixed effects and robust standard errors.')

# Stage coefficient and marginal-effect tables reuse the same source row components.
t=clone(3)
for row in list(t._tbl.findall(qn('w:tr')))[2:]:t._tbl.remove(row)
stage_terms=[('VD','VD'),('VD × Series A','1.VD#2.fundstage'),('VD × Series B','1.VD#3.fundstage'),
 ('VD × Series C+','1.VD#4.fundstage'),('VD × Prior Patent','1.VD#1.priorpatentdummy_VD'),
 ('VD × Early Stage','1.VD#1.Earlystage_VD'),('Series A','2.fundstage'),('Series B','3.fundstage'),
 ('Series C+','4.fundstage'),('Prior Patent','priorpatentdummy_VD'),('Early Stage','Earlystage_VD'),
 ('Log(Total Funding at VD)','ln_MoneyRaised_VD'),('Funding Rounds at VD','NumofFundingRounds_VD'),('Constant','_cons')]
# Split coefficient panel into two source-derived pages to avoid squeezing 28 rows.
for panel,items in [('A',stage_terms[:7]),('B',stage_terms[7:])]:
    st=Table(deepcopy(t._tbl),D._body);cell(st,0,0,'DV: Log(1 + Normalized Granted Applications)')
    for label,term in items:
        for ri in [2,3]:st._tbl.append(deepcopy(Table(templates[3],D._body).rows[ri]._tr))
        r=len(st.rows)-2;cell(st,r,0,label)
        for j,mid in enumerate([5,7,9,11],1):put(st,r,j,mid,term)
    add(f'Table 3-2. Funding-stage Heterogeneity — Panel {panel}',st,
        'Seed/pre-Series is the stage reference category. All columns include region, industry and founding-year fixed effects; N=2,820. '
        'Stage assignments include source imputations. Panel B continues the same four models. '+common)
st=Table(deepcopy(t._tbl),D._body);cell(st,0,0,'Average VD effect within funding stage')
for stage,label in enumerate(['Seed / pre-Series','Series A','Series B','Series C+'],1):
    for ri in [2,3]:st._tbl.append(deepcopy(Table(templates[3],D._body).rows[ri]._tr))
    r=len(st.rows)-2;cell(st,r,0,label)
    for j,mid in enumerate([6,8,10,12],1):
        term=('1bn' if stage==1 else str(stage))+'.fundstage'
        put(st,r,j,mid,term)
add('Table 3-2. Funding-stage Heterogeneity — Panel C',st,
    'Average discrete VD effects from margins fundstage, dydx(1.VD), corresponding to Panels A–B. '+common)
regtable('Table 4. VD Backing and Post-VD Patent Quality',range(13,17),'DV: Log(1 + Normalized Citations per Patent)',
         'DV: ln(1 + norm_cite_perpat_win3_v4). Undefined per-patent outcomes are excluded. OLS with robust standard errors and the listed fixed effects.')
regtable('Table 5. Cumulative Subsequent Financing',range(17,21),'DV: Log(1 + Total Subsequent Financing)',
         'DV: ln(1 + moneyraised_win3), all financing types in t+1 through t+3, not equity alone. OLS with robust standard errors.')
t=clone(6)
for j,mid in enumerate([21,22,23],1):
    for r,term in {3:'VD',5:'priorpatentdummy_VD',7:'Earlystage_VD',9:'ln_MoneyRaised_VD',11:'NumofFundingRounds_VD',13:'_cons'}.items():put(t,r,j,mid,term)
    z=C[C.model.eq(mid)].iloc[0];cell(t,19,j,num(z.N,True));cell(t,20,j,num(z.r2))
add('Table 6. Annual Subsequent Financing',t,
    'Columns (1)–(2) use all financing in t+1 and t+2; column (3) uses equity only in t+3, as coded. These are not a homogeneous financing series. '+common)

def cox6(title,models,extra,encoded=False):
    t=clone(7)
    for j,mid in enumerate(models,1):
        for r,term in {3:'VD',5:'priorpatentdummy_VD',7:'Earlystage_VD',9:'ln_MoneyRaised_VD',11:'NumofFundingRounds_VD'}.items():put(t,r,j,mid,term)
        z=C[C.model.eq(mid)].iloc[0];cell(t,17,j,num(z.fail,True));cell(t,18,j,num(z.N,True))
    if encoded:
        for r,label in [(14,'Region code control'),(15,'Industry code control'),(16,'Founding-year linear control')]:cell(t,r,0,label)
    add(title,t,'Cox log-hazard coefficients, not hazard ratios. Censoring year: 2026. '+extra+' '+common)
cox6('Table 7. Cox Proportional Hazard Models', [32,33,35,36,39,40],
     'Uses the later Table 7.v2 block with categorical fixed effects. IPO failure is explicitly reset before IPO estimation. Duration is stop year minus reference year.')
cox6('Table 7-2. Earlier Cox Specification',range(24,30),
     'The earlier code enters region, industry and founding-year codes as linear regressors, not categorical fixed effects. Duration is stop year minus reference year.',True)
t=clone(6)
for j,mid in enumerate([31,34,38],1):
    cell(t,1,j,['M&A','IPO','Defunct'][j-1])
    for r in [3,5,7,9,11,13]:
        for rr in [r,r+1]:cell(t,rr,j,'')
    put(t,3,j,mid,'VD')
    for r in [16,17,18]:cell(t,r,j,'No')
    z=C[C.model.eq(mid)].iloc[0];cell(t,19,j,num(z.N,True));cell(t,20,j,num(z.fail,True))
cell(t,0,0,'DV: Exit hazard');cell(t,20,0,'Number of failures')
add('Table 7-3. Unadjusted Cox Models',t,'VD-only models on the corresponding Table 7 analysis samples. Log-hazard coefficients. '+common)
t=clone(11)
for j,mid in enumerate([41,42],1):
    for r,term in {3:'ln_vdsupply',5:'VD',7:'priorpatentdummy_VD',9:'Earlystage_VD',11:'ln_MoneyRaised_VD',13:'NumofFundingRounds_VD'}.items():put(t,r,j,mid,term)
    z=C[C.model.eq(mid)].iloc[0]
    for r,key in [(19,'N'),(20,'F'),(21,'r2')]:cell(t,r,j,num(z[key],key=='N'))
    cell(t,22,j,num(z.kp) if j==1 else '')
add('Table A1. Instrumental-variable Estimates',t,'Column (1): robust 2SLS; column (2): first stage on the same sample. Instrument: ln(1 + regional VD supply). '+common)
regtable('Table A2. Patenting Startups',range(43,47),'DV: Log(1 + Normalized Granted Applications)',
         'Table 3 specifications restricted to a positive, nonmissing normalized granted-application outcome. Robust standard errors.',template=12)
for idx,start,outcome in [(13,48,'M&A'),(14,53,'IPO'),(15,58,'Defunct')]:
    regtable(f'Table A{idx-10}. Cox Models with Moderating Effects — {outcome}',range(start,start+4),
             'DV: Log Hazard of '+outcome,
             'Cox coefficients. Inclusive duration = stop year minus reference year plus one; censoring in 2026. This differs from Table 7. Categorical region, industry and founding-year controls.',template=idx,cox=True)
for number,start,dv,title,detail in [
 (8,74,'DV: Mean Gunning Fog','Gunning Fog Readability','OLS; robust standard errors.'),
 (9,78,'DV: Mean Automated Readability Index','Automated Readability Index','OLS; robust standard errors.'),
 (12,82,'DV: Log(1 + Existing-domain Patents)','Patents in Existing Technological Domains','PPML is applied to ln(1 + oldpat_win3), exactly as coded, not the raw count. Coefficients are log-link coefficients. Conventional R-squared is not reported.'),
 (13,86,'DV: Number of New CPC Classes','Entry into New CPC Classes','PPML with robust standard errors; log-link coefficients. Conventional R-squared is not reported.'),
 (14,90,'DV: Mean Backward-cited CPC3 Count','Backward Citation Technological Breadth','OLS; robust standard errors.'),
 (15,94,'DV: Mean Self Citations per Patent','Self Citations','PPML on selfcite_win3, not total self citations or a self-citation share. Log-link coefficients; conventional R-squared is not reported.'),
 (19,62,'DV: Share of Uncited Patents','Uncited Patents','Fractional logit with robust standard errors. Link coefficients, not marginal effects. Conventional R-squared is not reported.'),
 (20,66,'DV: Breakthrough Patent Share','Breakthrough Patents','OLS on breakthrough_win3 with robust standard errors.'),
 (21,70,'DV: Log(1 + Mean Claims)','Patent Claims','OLS on ln(1 + claims_win3) with robust standard errors.')]:
    regtable(f'Table A{number}. {title}',range(start,start+4),dv,
             detail+' Outcomes use the complete t+1 to t+3 filing window. Unavailable source information and undefined portfolios remain missing, not zero.',template=18)

# Replace retained reference row labels too (annual financing, Cox and IV templates).
for table in D.tables:
    for row in table.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for run in p.runs:
                    if 'No Prior Patent' in run.text:
                        run.text=run.text.replace('No Prior Patent','Prior Patent')
assert 'No Prior Patent' not in ''.join(D._element.itertext())
# Save only edited document.xml, preserving all other source package parts.
buf=BytesIO();D.save(buf)
with zipfile.ZipFile(buf) as zz:xml=zz.read('word/document.xml')
with zipfile.ZipFile(SOURCE) as src,zipfile.ZipFile(FINAL,'w',zipfile.ZIP_DEFLATED) as dst:
    for item in src.infolist():dst.writestr(item,xml if item.filename=='word/document.xml' else src.read(item.filename))
with zipfile.ZipFile(SOURCE) as src,zipfile.ZipFile(FINAL) as dst:
    assert all(src.read(n)==dst.read(n) for n in src.namelist() if n!='word/document.xml')
assert not list(OUT.glob('*.xls*'))
json.dump({'tables':inventory,'cells':audit,'reference_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest()},
          open(OUT/'word_numeric_audit.json','w'),indent=2)
print(FINAL);print('Tables',len(inventory),'coefficient cells',len(audit))
