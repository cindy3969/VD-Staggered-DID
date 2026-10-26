"""Patch only significance formatting in the existing table document."""
from pathlib import Path
from io import BytesIO
import zipfile, json
import pandas as pd
from docx import Document

root=Path(__file__).resolve().parents[1]
out=root/'outputs/table_update_20261019'
path=out/'VD Table 20261019.docx'
original=path.read_bytes()
d=Document(BytesIO(original))
c=pd.read_stata(out/'main_coefficients.dta')
i=pd.read_stata(root/'outputs/psm_final_2820_package/table3_application_results/innovation/innovation_models.dta')
assert len(d.tables)==29
untouched={k:d.tables[k]._tbl.xml for k in [0,1,2,8,9,10,16,17]}
checks=[]
def stars(p):
    return '***' if p<.01 else '**' if p<.05 else '*' if p<.1 else ''
def number(b):
    return f'{0 if abs(b)<.0005 else b:.3f}'
def pvalue(p):
    return '' if pd.isna(p) else f'({p:.3f})'
def setcell(t,r,col,text):
    cell=d.tables[t].cell(r,col)
    p=cell.paragraphs[0]
    if p.runs:
        p.runs[0].text=text
        for run in p.runs[1:]: run.text=''
    else: p.add_run(text)
    for p in cell.paragraphs[1:]:
        for run in p.runs: run.text=''
    checks.append((t,r,col,text))
terms={2:'VD',4:'1.VD#1.nopriorpatentdummy_VD',6:'1.VD#1.Earlystage_VD',8:'nopriorpatentdummy_VD',10:'Earlystage_VD',12:'ln_MoneyRaised_VD',14:'NumofFundingRounds_VD',16:'_cons'}
maps=[(3,[1,2,3,4],terms),(4,[5,6,7,8],terms),(5,[9,10,11,12],terms),
      (6,[13,14,15],{3:'VD',5:'nopriorpatentdummy_VD',7:'Earlystage_VD',9:'ln_MoneyRaised_VD',11:'NumofFundingRounds_VD',13:'_cons'}),
      (7,[24,25,27,28,31,32],{3:'VD',5:'nopriorpatentdummy_VD',7:'Earlystage_VD',9:'ln_MoneyRaised_VD',11:'NumofFundingRounds_VD'}),
      (11,[33,34],{3:'ln_vdsupply',5:'VD',7:'nopriorpatentdummy_VD',9:'Earlystage_VD',11:'ln_MoneyRaised_VD',13:'NumofFundingRounds_VD'}),
      (12,[35,36,37,38],terms)]
for t,models in [(13,[40,41,42,43]),(14,[45,46,47,48]),(15,[50,51,52,53])]:
    maps.append((t,models,{r:v for r,v in terms.items() if r!=16}))
for t,start in enumerate(range(54,86,4),18): maps.append((t,list(range(start,start+4)),terms))
for t,models,rows in maps:
    for col,m in enumerate(models,1):
        z=c[c.model.eq(m)]
        for r,term in rows.items():
            v=z[z.term.str.split(':').str[-1].isin([term,'1.'+term]) & z.se.gt(0)]
            if v.empty: continue
            assert len(v)==1
            v=v.iloc[0]
            old=d.tables[t].cell(r,col).text
            assert old in [number(v.b),number(v.b)+stars(v.p)],(t,r,col,old)
            setcell(t,r,col,number(v.b)+stars(v.p))
            setcell(t,r+1,col,pvalue(v.p))
dvs=['nocite_win3','breakthrough_win3','originality_win3','generality_win3','cite_diversity_win3','newshare_win3','own_dist_win3','selfshare_win3','partner_search_win3','desc_words_win3','cite_breadth_win3','selfcite_win3','partner_dist_win3','newpat_win3','oldpat_win3','newcpc_win3']
for row,dv in enumerate(dvs,1):
    for col in range(1,6):
        v=i[i.dv.eq(dv)&i.model.eq(f'ols{col}_VD_AME')].iloc[0]
        setcell(26,row,col,number(v.b)+stars(v.p)+'\n'+pvalue(v.p))
for row,v in enumerate(i[i.model.isin(['fraclogit_VD_AME','poisson_VD_AME'])].itertuples(),1):
    setcell(27,row,3,number(v.b)+stars(v.p)); setcell(27,row,4,pvalue(v.p))
for row,v in enumerate(i[i.dv.eq('new_minus_old')].itertuples(),1):
    setcell(28,row,2,number(v.b)+stars(v.p)); setcell(28,row,4,pvalue(v.p))
legend=' P-values are rounded to three decimal places. * p<0.10, ** p<0.05, *** p<0.01; stars use unrounded p-values.'
targets={9,12,16,21,24,39,43,47,52,56}
for idx,p in enumerate(d.paragraphs):
    text=p.text
    text=text.replace('; values below 0.001 are shown as <0.001','')
    if idx in targets or (idx>60 and text.startswith('Note:')):
        if legend.strip() not in text: text+=legend
    if text!=p.text:
        p.runs[0].text=text
        for run in p.runs[1:]: run.text=''
assert '<0.001' not in d._element.xml
buffer=BytesIO()
with zipfile.ZipFile(BytesIO(original)) as src, zipfile.ZipFile(buffer,'w',zipfile.ZIP_DEFLATED) as dst:
    for entry in src.infolist():
        dst.writestr(entry,d._element.xml.encode('utf-8') if entry.filename=='word/document.xml' else src.read(entry.filename))
result=buffer.getvalue()
test=Document(BytesIO(result))
for t,r,col,value in checks: assert test.tables[t].cell(r,col).text==value
for t,xml in untouched.items(): assert test.tables[t]._tbl.xml==xml
with zipfile.ZipFile(BytesIO(original)) as a,zipfile.ZipFile(BytesIO(result)) as b:
    for name in a.namelist():
        if name!='word/document.xml': assert a.read(name)==b.read(name)
assert path.read_bytes()==original,'File changed during editing'
path.write_bytes(result)
(out/'significance_format_audit.json').write_text(json.dumps({'checked_cells':len(checks),'star_thresholds':[.1,.05,.01],'p_decimals':3,'excluded_tables_preserved':True,'layout_verified':False},indent=2),encoding='utf-8')
print('Updated existing file:',path)
print('Verified cells:',len(checks))
