from pathlib import Path
from zipfile import ZipFile
import shutil, json, re
from math import erfc, sqrt
from copy import deepcopy
import pandas as pd
from lxml import etree as E
from docx import Document

ROOT=Path('C:/Users/user/Documents/VD')
SRC=Path('C:/Users/user/iCloudDrive/5학기/연구미팅/VD PPT/20261022 VD Table.docx')
OUT=ROOT/'outputs/refresh_did_20261023'
BACK=OUT/'before_tables9_14.docx'
if not BACK.exists(): shutil.copy2(SRC,BACK)
d=Document(SRC)
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
original=deepcopy(d._element)
changed=set()
audit=[]
def text(el): return ''.join(el.xpath('.//w:t/text()'))
def put(t,r,c,value):
    el=d.tables[t].rows[r]._tr.tc_lst[c]
    old=text(el)
    if old==value and '\n' not in value: return
    chunks=value.split('\n')
    containers=el.xpath('./w:p') if len(chunks)>1 else [el]
    assert len(containers)==len(chunks)
    for container,rest in zip(containers,chunks):
        nodes=container.xpath('.//w:t')
        if not nodes:
            assert not rest, (t,r,c,value)
            continue
        for n in nodes[:-1]:
            k=len(n.text or '')
            n.text,rest=rest[:k],rest[k:]
        nodes[-1].text=rest
    changed.add(t); audit.append([t,r,c,old,value])
def stars(p): return '***' if p<.01 else '**' if p<.05 else '*' if p<.1 else ''
def coef(x,n=3): return f'{x.att:.{n}f}'+stars(x.p)
def pick(df,s,agg,e):
    a=df[(df['sample']==s)&(df.aggregation==agg)&(df.effect==e)]
    assert len(a)==1,(s,agg,e,a)
    return a.iloc[0]
def load_results(dv,groups):
    log=(OUT/(dv+'.log')).read_text(errors='replace')
    if 'END_SPEC v4 all' not in log:return None
    block=log.split('BEGIN_SPEC v4 all ')[1].split('END_SPEC v4 all')[0]
    N,firms=map(int,re.search(r'N=(\d+) firms=(\d+)',block).groups())
    rows=[]
    for e,b,se in re.findall(r'^\s*(ATT|Pre_avg|Post_avg|Tm[1-4]|Tp[0-4]|GAverage)\s*\|\s*([-+.\deE]+)\s+([-+.\deE]+)',block,re.M):
        b,se=float(b),float(se)
        p=erfc(abs(b/se)/sqrt(2))
        assert min(abs(p-v) for v in [.01,.05,.1])>1e-5
        rows.append(dict(sample='all',aggregation='simple' if e=='ATT' else 'group' if e=='GAverage' else 'event',effect=e,att=b,se=se,p=p,N=N,firms=firms,rc=0))
    assert len(rows)==13
    frames=[pd.DataFrame(rows)]
    for s in groups:
        if s=='all':continue
        f=OUT/(dv+'_'+s+'_results.dta');l=OUT/(dv+'_'+s+'.log')
        if not l.exists() or 'REQUESTED_DID_COMPLETE' not in l.read_text(errors='replace'):return None
        frames.append(pd.read_stata(f,convert_categoricals=False))
    return pd.concat(frames,ignore_index=True)
for dv,t,groups in [('grantedapp_norm',8,['all','noprior','prior','nonlate','late']),('norm_cite_perpat_v4',10,['all','nonlate','late'])]:
    df=load_results(dv,groups+['notyet'])
    if df is None:continue
    assert df.rc.eq(0).all()
    for c,s in enumerate(groups,1):
        a=pick(df,s,'simple','ATT'); b=pick(df,s,'event','Post_avg')
        for r,x in [(1,a),(3,b)]:
            put(t,r,c,coef(x));put(t,r+1,c,f'({x.p:.3f})')
        put(t,9,c,f'{int(a.N):,}');put(t,10,c,f'{int(a.firms):,}')
        if s=='nonlate': put(t,0,c,f'({c})\nSeed–Series B')
        if s=='late': put(t,0,c,f'({c})\nSeries C+')
    if t==10: put(t,0,0,'DV: Patent citations')
    for k,s in enumerate(['all','notyet']):
        tt=t+1
        def cellval(x): return coef(x,2)+f' ({x.se:.2f})'
        put(tt,1,k+1,cellval(pick(df,s,'simple','ATT')))
        for r,e in enumerate(['Tm4','Tm3','Tm2','Tm1','Tp0','Tp1','Tp2','Tp3','Tp4'],3):
            put(tt,r,2+2*k,cellval(pick(df,s,'event',e)))
        put(tt,12,2+2*k,cellval(pick(df,s,'event','Post_avg')))
        put(tt,13,2+2*k,cellval(pick(df,s,'group','GAverage')))

tw=ROOT/'outputs/vd_intensity/tables_3_1_to_3_7'
if 'Finished:' in (tw/'replication.log').read_text(errors='replace'):
    df=pd.read_stata(tw/'regression_results.dta',convert_categoricals=False)
    assert 'lag_late' in set(df.term)
    terms=['dose_post','dose_post_prior','dose_post_late','post_prior','post_late','post','prior_patent_l1','lag_late','ln1p_cum_funding_l1','ln1p_rounds_l1']
    for legacy,t in [(1,12),(4,13)]:
        for m in range(1,5):
            model=df[(df.table==legacy)&(df.model==m)]
            for j,term in enumerate(terms):
                a=model[model.term==term]
                if len(a):
                    x=a.iloc[0];put(t,2*j+1,m,f'{x.b:.3f}'+stars(x.p));put(t,2*j+2,m,f'({x.p:.3f})')
            x=model.iloc[0]
            put(t,24,m,f'{int(x.N):,}');put(t,25,m,f'{int(x.firms):,}');put(t,26,m,f'{x.within_r2:.3f}')
        put(t,15,0,'Lagged late stage')
        # Columns are nested interaction specifications, not subsamples.
        for m,label in enumerate(['All firms','Prior int.','Late int.','Both int.'],1):put(t,0,m,f'({m})\n{label}')

# Every formatting element and all non-target tables remain intact.
for tag in ['tblPr','tblGrid','trPr','tcPr','pPr','rPr','sectPr']:
    a=original.xpath('//w:'+tag);b=d._element.xpath('//w:'+tag)
    assert [E.tostring(x) for x in a]==[E.tostring(x) for x in b],tag
for i,t in enumerate(d.tables):
    if i not in changed:
        assert E.tostring(original.xpath('//w:body/w:tbl')[i])==E.tostring(t._tbl),i
dest=OUT/'20261022 VD Table.docx'
with ZipFile(SRC) as zin,ZipFile(dest,'w') as zout:
    for z in zin.infolist():
        data=E.tostring(d._element,xml_declaration=True,encoding='UTF-8',standalone=True) if z.filename=='word/document.xml' else zin.read(z.filename)
        zout.writestr(z,data)
with ZipFile(SRC) as a,ZipFile(dest) as b:
    assert all(a.read(n)==b.read(n) for n in a.namelist() if n!='word/document.xml')
(OUT/'word_update_audit.json').write_text(json.dumps({'tables':sorted(changed),'cells':audit},ensure_ascii=False,indent=2),encoding='utf-8')
print('Updated physical tables',sorted(changed),'cells',len(audit),dest)
