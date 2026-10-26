from pathlib import Path
import json
import openpyxl

root = Path(r"C:\Users\user\iCloudDrive\5학기\연구미팅\20240108-2010 to 2015 firms")
files = list(root.glob("*.xlsx")) + list((root / "python").glob("*.xlsx"))
out = []
for path in files:
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            header_rows = list(sheet.iter_rows(min_row=1, max_row=min(5, sheet.max_row), values_only=True))
            flattened = [str(value) if value is not None else "" for row in header_rows for value in row][:100]
            record = {
                "file": str(path), "sheet": sheet.title,
                "rows": sheet.max_row, "cols": sheet.max_column,
                "header": flattened,
            }
            text = " ".join(flattened).lower()
            if (2000 <= sheet.max_row <= 2100 or 18000 <= sheet.max_row <= 18500
                    or any(key in text for key in ["citation", "patent", "cpc", "filing"])):
                out.append(record)
        workbook.close()
    except Exception as exc:
        out.append({"file": str(path), "error": str(exc)})
print(json.dumps(out, ensure_ascii=False, indent=2))
