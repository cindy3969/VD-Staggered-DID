"""Update the retained table document using freshly captured Stata estimates.
Only word/document.xml is replaced; other original package parts are preserved.
"""
from pathlib import Path
from copy import deepcopy
from io import BytesIO
import hashlib, json, zipfile
from datetime import datetime
import numpy as np
import pandas as pd
from statistics import NormalDist
from docx import Document
from docx.table import Table
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/table_update_20261019'
SOURCE=Path(r'C:\Users\user\iCloudDrive\5학기\연구미팅\VD PPT\20260822 Table.docx')
FINAL=OUT/'VD Table 20261019.docx'
PKG=ROOT/'outputs/psm_final_2820_package'
D=Document(SOURCE); P=list(D.paragraphs); T=list(D.tables)
C=pd.read_csv(OUT/'main_coefficients.csv')
I=pd.read_csv(PKG/'table3_application_results/innovation/innovation_models.csv')
X=pd.read_stata(PKG/'table3_current_2820_cross_section_v2.dta',convert_categoricals=False)
matched=pd.read_stata(PKG/'psm_stata_matched_2820.dta',convert_categoricals=False)
assert X.OrganizationName.is_unique and matched.OrganizationName.is_unique
assert set(X.OrganizationName)==set(matched.OrganizationName)
assert X.VD.value_counts().to_dict()=={0:1410,1:1410}
for name in ['RUN_TABLE3_PATENT_APPLICATION_WIN3','RUN_INNOVATION_WIN3']:
    assert (PKG/(name+'.do')).read_bytes()==(OUT/(name+'_source_snapshot.do')).read_bytes()
assert C.model.nunique()==85
assert len(I)==93 and I.rc.eq(0).all() and I[['b','p','se']].notna().all().all()
edits=[]

def corr_star(a,b):
    r=float(np.corrcoef(a,b)[0,1]); df=len(a)-2
    # Student-t 0.975 quantile expansion; df=2818 for complete matched data.
    # O(df^-4) error is negligible here; assert no result near the boundary.
    z=NormalDist().inv_cdf(.975)
    critical=z+(z**3+z)/(4*df)+(5*z**5+16*z**3+3*z)/(96*df**2)+(3*z**7+19*z**5+17*z**3-15*z)/(384*df**3)
    t=abs(r)*np.sqrt(df/(1-r*r))
    assert abs(t-critical)>1e-6
    return r,t>critical

def par(p,text):
    if p.runs:
        p.runs[0].text=text
        for r in p.runs[1:]: r.text=''
    else: p.add_run(text)

def cell(t,r,c,text):
    target=t.cell(r,c); par(target.paragraphs[0],str(text))
    for p in target.paragraphs[1:]: par(p,'')

def num(x,integer=False):
    if pd.isna(x): return ''
    if integer: return f'{int(x):,}'
    if abs(x)<0.0005: x=0
    return f'{x:.3f}'

def pv(x):
    if pd.isna(x): return ''
    return '(<0.001)' if x<0.001 else f'({x:.3f})'

def coef(mid,term):
    z=C[C.model.eq(mid)].copy(); z['clean']=z.term.str.split(':').str[-1]
    candidates=[term,'1.'+term]
    z=z[z.clean.isin(candidates)]
    z=z[z.se.gt(0)]
    assert len(z)<=1,(mid,term,z.to_dict())
    return z.iloc[0] if len(z) else None

def putcoef(t,r,c,mid,term):
    v=coef(mid,term); b=num(v.b) if v is not None else ''; p=pv(v.p) if v is not None else ''
    cell(t,r,c,b); cell(t,r+1,c,p)
    edits.append(dict(model=mid,term=term,coefficient=b,p=p))

TERMS={2:'VD',4:'1.VD#1.nopriorpatentdummy_VD',6:'1.VD#1.Earlystage_VD',
       8:'nopriorpatentdummy_VD',10:'Earlystage_VD',12:'ln_MoneyRaised_VD',14:'NumofFundingRounds_VD',16:'_cons'}
LABELS={2:'VD',4:'VD x No Prior Patent',6:'VD x Early Stage',8:'No Prior Patent',
        10:'Early Stage',12:'Log(Total Funding at VD)',14:'Funding Rounds at VD',16:'Constant'}

def fill_reg(t,models,dv,cox=False):
    cell(t,0,0,dv)
    for r,term in TERMS.items():
        if cox and r==16: continue
        cell(t,r,0,LABELS[r])
        for col,mid in enumerate(models,1): putcoef(t,r,col,mid,term)
    for j,mid in enumerate(models,1):
        z=C[C.model.eq(mid)].iloc[0]
        cell(t,21 if cox else 22,j,num(z.N,True))
        cell(t,20 if cox else 23,j,num(z.fail,True) if cox else num(z.r2))
    for row in t.rows:
        if row.cells[0].text.strip()=='Year FE': par(row.cells[0].paragraphs[0],'Founding-year FE')

# Table 1 and two panels of Table 2: current do-file variables and Stata output.
S=pd.read_stata(OUT/'table1_stats.dta').set_index('variable')
R=pd.read_stata(OUT/'table2_correlations.dta').set_index(['variable1','variable2'])
statvars=['patent_application_win3','VD','nopriorpatentdummy_VD','Earlystage_VD',
          'MoneyRaised_VD','NumofFundingRounds_VD','ind_1','ind_2','ind_4','reg_1','reg_2','reg_3']
stats=pd.DataFrame({
 'Patent Applications (3-year filing window)':X.patent_application_win3,
 'VD':X.VD,'No Prior Patent':1-X.priorpatentdummy_VD,'Early Stage':X.earlystage,
 'Total Funding at VD (USD millions)':X.MoneyRaised_VD/1e6,'Funding Rounds at VD':X.NumofFundingRounds_VD,
 'Bio/Healthcare':X.Industry2.eq(2).astype(int),'Hardware/Manufacturing/Deep Tech':X.Industry2.eq(3).astype(int),
 'Software/IT services':X.Industry2.eq(5).astype(int),'California':X.Region2.eq(1).astype(int),
 'Massachusetts':X.Region2.eq(2).astype(int),'New York':X.Region2.eq(3).astype(int)})
for row,(name,s) in enumerate(stats.items(),1):
    cell(T[0],row,0,name)
    values=S.loc[statvars[row-1],['N','mean','sd','min','max']].to_numpy(float,copy=True)
    if row==5: values[1:]/=1e6
    assert np.allclose(values,[s.count(),s.mean(),s.std(),s.min(),s.max()])
    for col,value in enumerate(values,1): cell(T[0],row,col,num(value,col==1))
for row in range(1,13):
    cell(T[1],row,1,stats.columns[row-1])
    for col in range(2,8):
        j=col-2; i=row-1
        if i<j: value=''
        elif i==j: value='1.000'
        else:
            result=R.loc[(statvars[i],statvars[j])]; value=num(result.r)+('*' if result.p<.05 else '')
        cell(T[1],row,col,value)
for row in range(1,7):
    i=row+5; cell(T[2],row,1,stats.columns[i])
    for col in range(2,8):
        j=col+4
        if i<j: value=''
        elif i==j: value='1.000'
        else:
            result=R.loc[(statvars[i],statvars[j])]; value=num(result.r)+('*' if result.p<.05 else '')
        cell(T[2],row,col,value)

fill_reg(T[3],[1,2,3,4],'DV: Log(1 + Cohort-normalized Granted Applications)')
fill_reg(T[4],[5,6,7,8],'DV: Log(1 + Normalized Citations per Patent)')
fill_reg(T[5],[9,10,11,12],'DV: Log(1 + Total Subsequent Financing)')
for j,mid in enumerate([13,14,15],1):
    for r,term in {3:'VD',5:'nopriorpatentdummy_VD',7:'Earlystage_VD',9:'ln_MoneyRaised_VD',11:'NumofFundingRounds_VD',13:'_cons'}.items(): putcoef(T[6],r,j,mid,term)
    z=C[C.model.eq(mid)].iloc[0]; cell(T[6],19,j,num(z.N,True)); cell(T[6],20,j,num(z.r2))
cell(T[6],0,0,'DV: Log(1 + Annual Subsequent Financing)')
cell(T[6],1,1,'t+1 All financing'); cell(T[6],1,2,'t+2 All financing'); cell(T[6],1,3,'t+3 Equity only')
cell(T[6],5,0,'No Prior Patent'); cell(T[6],18,0,'Founding-year FE')
# Use the later Table 7.v2 specification, with categorical fixed effects.
for j,mid in enumerate([24,25,27,28,31,32],1):
    for r,term in {3:'VD',5:'nopriorpatentdummy_VD',7:'Earlystage_VD',9:'ln_MoneyRaised_VD',11:'NumofFundingRounds_VD'}.items(): putcoef(T[7],r,j,mid,term)
    z=C[C.model.eq(mid)].iloc[0]; cell(T[7],17,j,num(z.fail,True)); cell(T[7],18,j,num(z.N,True))
cell(T[7],5,0,'No Prior Patent'); cell(T[7],16,0,'Founding-year FE')

# IV table: distinct first/second-stage diagnostics, never duplicate a statistic.
for j,mid in enumerate([33,34],1):
    for r,term in {3:'ln_vdsupply',5:'VD',7:'nopriorpatentdummy_VD',9:'Earlystage_VD',11:'ln_MoneyRaised_VD',13:'NumofFundingRounds_VD'}.items(): putcoef(T[11],r,j,mid,term)
    z=C[C.model.eq(mid)].iloc[0]
    cell(T[11],19,j,num(z.N,True)); cell(T[11],20,j,num(z.F)); cell(T[11],21,j,num(z.r2))
cell(T[11],1,1,'Log(1 + Cohort-normalized Granted Applications)')
cell(T[11],7,0,'No Prior Patent'); cell(T[11],18,0,'Founding-year FE')
cell(T[11],22,0,'Kleibergen-Paap rk Wald F')
cell(T[11],22,1,''); cell(T[11],22,2,num(C[C.model.eq(33)].iloc[0].kp))
fill_reg(T[12],[35,36,37,38],'DV: Log(1 + Cohort-normalized Granted Applications)')
fill_reg(T[13],[40,41,42,43],'DV: Log Hazard of M&A Exit',True)
fill_reg(T[14],[45,46,47,48],'DV: Log Hazard of IPO Exit',True)
fill_reg(T[15],[50,51,52,53],'DV: Log Hazard of Defunct Outcome',True)

common='No Prior Patent equals one minus the prior-patent indicator. Early Stage and funding controls are measured at the VD/reference year. All models include region, industry and founding-year fixed effects. Two-sided p-values are in parentheses; values below 0.001 are shown as <0.001.'
par(P[1],'Note: Summary statistics for 2,820 rematched firms (1,410 VD and 1,410 non-VD), updated 19 October 2026. Patent applications are patent_application_win3, as specified in the analysis do-file. Total funding is reported in USD millions, without a log transformation. No Prior Patent is the complement of the prior-patent indicator. Indicator means are proportions.')
par(P[9],'Note: The DV is ln(1 + cohort-normalized granted applications filed in t+1 through t+3), using grantedpatentbyfiling_norm_win3. This differs from an unnormalized patent count. '+common+' Heteroskedasticity-robust standard errors are used.')
par(P[12],'Note: The DV is ln(1 + normalized citations per patent), using norm_cite_perpat_win3_v4 for the three-year post-reference filing window. Firms without a defined per-patent quality measure are excluded, not assigned zero quality. '+common)
par(P[15],'TABLE 5. Cumulative Subsequent Financing after VD Backing')
par(P[16],'Note: The DV is ln(1 + moneyraised_win3), encompassing all financing types in t+1 through t+3, not equity alone. Missing or undisclosed financing is excluded. '+common)
par(P[20],'TABLE 6. Annual Subsequent Financing after VD Backing')
par(P[21],'Note: Following the executed specification, columns (1)-(2) use all financing in t+1 and t+2, while column (3) uses equity financing in t+3. These are not a homogeneous financing series. Missing outcomes are excluded separately by horizon. '+common)
par(P[24],'Note: Cox coefficients are log-hazard coefficients, not hazard ratios. The later Table 7.v2 specification is used: columns (1), (3), and (5) include categorical region, industry and founding-year fixed effects; the paired columns add the listed controls. Duration is exit/censoring year minus reference year, with censoring in 2026 and nonpositive durations excluded. The IPO risk set is explicitly reset to IPO failures. p-values are in parentheses.')
for idx in [28,32,35]: par(P[idx],P[idx].text+' Retained from the August 2026 document; not re-estimated in the 19 October 2026 update.')
par(P[39],'Note: Column (1) is robust 2SLS and column (2) is the first stage estimated on the same sample. The instrument is ln(1 + regional VD supply). F-stat reports the respective model F statistic; the final row reports the heteroskedasticity-robust Kleibergen-Paap rk Wald F from 2SLS. '+common)
par(P[43],'Note: Table 3 is re-estimated among firms with a positive, nonmissing cohort-normalized granted-application outcome in t+1 through t+3. Sample sizes and fit statistics reflect the current estimation. '+common)
for idx,outcome in [(47,'M&A'),(52,'IPO'),(56,'defunct')]:
    par(P[idx],f'Note: Cox log-hazard coefficients for {outcome} with moderator interactions. The appendix uses its own inclusive duration definition: exit/censoring year minus reference year plus one, censored in 2026. This differs from Table 7. '+common)
par(P[60],P[60].text+' Matching diagnostics and the associated density figures are retained from the source document; the two regression scripts do not regenerate them.')

# Append source-style tables without altering existing employee/revenue tables.
def append_caption(title):
    p=D.add_paragraph(); par(p,title); p.paragraph_format.page_break_before=True
    p.paragraph_format.keep_with_next=True; p.paragraph_format.space_after=Pt(6)
    for r in p.runs: r.font.name='Times New Roman'; r.font.size=Pt(11); r.bold=True; r.font.color.rgb=RGBColor(0,0,0)
    return p

def append_note(text):
    p=D.add_paragraph(text); p.paragraph_format.space_before=Pt(6); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1
    for r in p.runs: r.font.name='Times New Roman'; r.font.size=Pt(10)

def prepare_new_table(t,widths):
    t.autofit=False
    for col,w in zip(t.columns,widths): col.width=Inches(w)
    for row in t.rows:
        for c,w in zip(row.cells,widths):
            c.width=Inches(w)
            for p in c.paragraphs:
                p.paragraph_format.line_spacing=1; p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0)
                for r in p.runs: r.font.name='Times New Roman'; r.font.size=Pt(10)
        trpr=row._tr.get_or_add_trPr()
        for el in list(trpr):
            if el.tag==qn('w:trHeight'): trpr.remove(el)
        if not trpr.find(qn('w:cantSplit')) is not None: trpr.append(OxmlElement('w:cantSplit'))
    head=t.rows[0]._tr.get_or_add_trPr(); head.append(OxmlElement('w:tblHeader'))

new_groups=[
 ('A8','Gunning Fog Readability',[54,55,56,57],'DV: Mean Gunning Fog','Higher values indicate less readable patent descriptions.'),
 ('A9','Automated Readability Index',[58,59,60,61],'DV: Mean Automated Readability Index','Higher values indicate a higher reading grade level.'),
 ('A10','Negative Flesch Reading Ease',[62,63,64,65],'DV: Mean Negative Flesch Reading Ease','Flesch reading ease is multiplied by -1; higher values indicate less readability.'),
 ('A11','Patents in New Technological Domains',[66,67,68,69],'DV: Log(1 + New-domain Patents)','New domains are primary CPC3 classes absent from the fixed pre-reference granted portfolio.'),
 ('A12','Patents in Existing Technological Domains',[70,71,72,73],'DV: Log(1 + Existing-domain Patents)','Existing domains are primary CPC3 classes present in the fixed pre-reference granted portfolio.'),
 ('A13','Entry into New CPC Classes',[74,75,76,77],'DV: Number of New Primary CPC3 Classes','This specification uses the class count in levels, not its logarithm.'),
 ('A14','Backward Citation Technological Breadth',[78,79,80,81],'DV: Mean Backward-cited CPC3 Count','This specification uses mean breadth in levels. It is not own-patent CPC4 invention scope.'),
 ('A15','Total Self Citations',[82,83,84,85],'DV: Log(1 + Total Self Citations)','The DV is the log of one plus the total backward self-citation count, not the self-citation share.')]
for tag,title,models,dv,definition in new_groups:
    append_caption(f'TABLE {tag}. {title}')
    xml=deepcopy(T[3]._tbl); D._body._body.insert(len(D._body._body)-1,xml); tab=Table(xml,D._body)
    fill_reg(tab,models,dv); prepare_new_table(tab,[2.65,.9,.9,.9,.9])
    append_note('Note: '+definition+' Outcomes use the t+1 to t+3 filing window. '+common+' These columns reproduce the main do-file; no model is selected based on significance.')

names={
 'nocite_win3':'No-citation share','breakthrough_win3':'Breakthrough share',
 'originality_win3':'Originality (own class excluded)','generality_win3':'Generality',
 'cite_diversity_win3':'Backward citation diversity','newshare_win3':'New-domain patent share',
 'own_dist_win3':'Distance from own portfolio','selfshare_win3':'Self-citation share',
 'partner_search_win3':'Partner search (already logged)','desc_words_win3':'Description words [log]',
 'cite_breadth_win3':'Citation breadth [log]','selfcite_win3':'Mean self-citations [log]',
 'partner_dist_win3':'Partner distance [log]','newpat_win3':'New-domain patents [log]',
 'oldpat_win3':'Existing-domain patents [log]','newcpc_win3':'New CPC3 classes [log]'}

def summary_table(headers,rows,widths):
    tab=D.add_table(rows=1,cols=len(headers))
    # Copy the reference regression table's border/look properties.
    pr=deepcopy(T[3]._tbl.tblPr); tab._tbl.replace(tab._tbl.tblPr,pr)
    for j,h in enumerate(headers): cell(tab,0,j,h)
    for row in rows:
        cells=tab.add_row().cells
        for c,v in zip(cells,row): par(c.paragraphs[0],str(v))
    prepare_new_table(tab,widths)
    for i,row in enumerate(tab.rows):
        for j,c in enumerate(row.cells):
            for p in c.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs: r.bold=i==0
    return tab

append_caption('TABLE A16. Additional Innovation Outcomes and Average VD Effects')
rows=[]
for dv,name in names.items():
    z=I[I.dv.eq(dv)].set_index('model'); vals=[]
    for s in range(1,6):
        r=z.loc[f'ols{s}_VD_AME']; vals.append(num(r.b)+'\n'+pv(r.p))
    rows.append([name,*vals,num(z.loc['ols1_VD_AME'].N,True)])
summary_table(['Outcome','(1)','(2)','(3)','(4)','(5)','N'],rows,[2.15,.7,.7,.7,.7,.7,.6])
append_note('Note: Average discrete VD effects from the separate innovation module. Columns (1)-(4) use the baseline, VD x No Prior Patent, VD x Early Stage, and both-interaction specifications. Column (5) adds reference-year fixed effects to column (1). Explicit fixed-effect indicators are used. [log] denotes ln(1 + DV). Only complete windows through the 2024 cutoff are included. Partner portfolios include usable VC/CVC/VD investors; own-portfolio comparisons require a nonempty baseline. p-values are in parentheses. These estimates are average VD effects, not the moderator-reference-category VD coefficients in Tables A8-A15.')

append_caption('TABLE A17. Fractional Response and Count Model Checks')
rows=[]
for r in I[I.model.isin(['fraclogit_VD_AME','poisson_VD_AME'])].itertuples():
    rows.append([names[r.dv].replace(' [log]',''),'Fractional logit' if r.model.startswith('frac') else 'Poisson',num(r.N,True),num(r.b),pv(r.p),f'{num(r.lo)}, {num(r.hi)}'])
summary_table(['Outcome','Model','N','VD effect','p-value','95% CI'],rows,[2.05,1,.55,.7,.65,1.3])
append_note('Note: Average discrete VD effects on the original DV scale, with robust standard errors. Fractional logit retains outcomes equal to zero or one. Poisson models include zero counts among baseline-eligible firms. All models include the baseline controls and region, industry and founding-year fixed effects. Missing portfolios or undefined ratios are not set to zero.')

append_caption('TABLE A18. Direct Tests of New versus Existing Domain Effects')
rows=[]
for r in I[I.dv.eq('new_minus_old')].itertuples():
    label='Stacked Poisson log IRR difference' if r.model.startswith('PPML') else 'Difference in log-count effects'
    rows.append([label,num(r.N,True),num(r.b),num(r.se),pv(r.p),f'{num(r.lo)}, {num(r.hi)}'])
summary_table(['Test','Firms','Difference','SE','p-value','95% CI'],rows,[2.05,.55,.8,.65,.65,1.55])
append_note('Note: Both tests compare the same firms and fixed pre-reference primary CPC3 portfolios. The log-count test uses ln(1 + new patents) minus ln(1 + existing patents). Stacked Poisson permits all covariate effects to vary by domain and clusters standard errors by firm. Negative differences indicate a more negative VD association for new domains. Neither two-sided test rejects equality at 5%; separate within-domain significance does not establish a differential effect. All patent-based measures retain the 2024 observation cutoff.')

# Save only the edited document XML into the original package.
payload=D._element.xml.encode('utf-8')
with zipfile.ZipFile(SOURCE) as src, zipfile.ZipFile(FINAL,'w',zipfile.ZIP_DEFLATED) as dst:
    for item in src.infolist(): dst.writestr(item,payload if item.filename=='word/document.xml' else src.read(item.filename))
check=Document(FINAL)
assert len(check.tables)==29
# Employee/revenue and matching tables are untouched, including formatting.
for idx in [8,9,10,16,17]: assert check.tables[idx]._tbl.xml==T[idx]._tbl.xml
with zipfile.ZipFile(SOURCE) as a, zipfile.ZipFile(FINAL) as b:
    for name in a.namelist():
        if name!='word/document.xml': assert a.read(name)==b.read(name),name
(OUT/'document_update_audit.json').write_text(json.dumps({'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
    'updated_at':datetime.now().isoformat(),
    'analysis_data_sha256':hashlib.sha256((PKG/'table3_current_2820_cross_section_v2.dta').read_bytes()).hexdigest(),
    'rematched_firm_ids_verified':True,'matched_firms':2820,'treated':1410,'controls':1410,
    'table1_table2_from_stata':True,'table1_funding_display_unit':'USD millions; no log',
    'tables_total':len(check.tables),'new_tables':'A8-A18','main_models':85,'innovation_models':93,
    'unreestimated_tables':['8','9','10','A6','A7'],'source_do_files_unchanged':True,
    'execution_copy_repairs':['Earlystage renamed reference -> Earlystage_VD','Reset IPO stset in Table 7.v2'],
    'table7_models':[24,25,27,28,31,32],'coefficient_cell_audit':edits},indent=2),encoding='utf-8')
print(FINAL)
print('Tables',len(check.tables),'coefficient pairs',len(edits))
