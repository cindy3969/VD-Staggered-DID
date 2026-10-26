from pathlib import Path
import re,json,sys,hashlib,zipfile
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/table_refresh_20261022'
OUT.mkdir(exist_ok=True)
SRC=ROOT/'outputs/psm_final_2820_package/RUN_TABLE3_PATENT_APPLICATION_WIN3_v3.do'
DOC=Path('C:/Users/user/iCloudDrive/5학기/연구미팅/VD PPT/20261020 VD Table.docx')
FINAL=OUT/'20261022 VD Table.docx'

def prepare():
    source=SRC.read_text(encoding='utf-8-sig')
    (OUT/'source_snapshot.do').write_text(source,encoding='utf-8')
    helper=r'''
global capture_id 0
postfile wordcapture int model str120 term str24 cmd double b se p N r2 r2p F df fail kp using "$OUT/coefficients.dta", replace
program define word_capture
 global capture_id = $capture_id + 1
 tempname B V bb ss pp
 matrix `B'=e(b)
 matrix `V'=e(V)
 local names : colfullnames `B'
 forvalues j=1/`=colsof(`B')' {
  local term : word `j' of `names'
  scalar `bb'=`B'[1,`j']
  scalar `ss'=sqrt(`V'[`j',`j'])
  scalar `pp'=.
  if `ss'>0 & !missing(`ss') {
   if !missing(e(df_r)) & e(cmd)!="ppmlhdfe" scalar `pp'=2*ttail(e(df_r),abs(`bb'/`ss'))
   else scalar `pp'=2*normal(-abs(`bb'/`ss'))
  }
  post wordcapture ($capture_id) ("`term'") ("`e(cmd)'") (`bb') (`ss') (`pp') (e(N)) (e(r2)) (e(r2_p)) (e(F)) (e(df_r)) (e(N_fail)) (e(widstat))
 }
 display "CAPTURED_MODEL " $capture_id " `e(cmd)'"
end
'''
    stats=r'''
local tablevars patent_application_win3 VD priorpatentdummy_VD latestage_VD MoneyRaised_VD NumofFundingRounds_VD ind_1 ind_2 ind_4 reg_1 reg_2 reg_3
postfile sumcapture str40 variable double N mean sd min max using "$OUT/summary.dta", replace
foreach v of local tablevars {
 quietly summarize `v'
 post sumcapture ("`v'") (r(N)) (r(mean)) (r(sd)) (r(min)) (r(max))
}
postclose sumcapture
postfile corcapture str40 variable1 str40 variable2 double N r p using "$OUT/correlations.dta", replace
foreach v of local tablevars {
 foreach w of local tablevars {
  quietly correlate `v' `w'
  local n=r(N)
  local rho=r(rho)
  local p=.
  if abs(`rho')<1 local p=2*ttail(`n'-2,abs(`rho')*sqrt((`n'-2)/(1-`rho'^2)))
  post corcapture ("`v'") ("`w'") (`n') (`rho') (`p')
 }
}
postclose corcapture
'''
    lines=source.splitlines(); out=[]; meta=[]; section=''; i=0
    while i<len(lines):
        start=i+1; group=[lines[i]];i+=1
        while group[-1].rstrip().endswith('///') and i<len(lines):
            group.append(lines[i]);i+=1
        cmd='\n'.join(group); s=cmd.strip()
        if s.startswith('*'):
            if re.search(r'Table\s*\w|A3 Pan',s,re.I):section=s.splitlines()[0]
            out.append(cmd);continue
        if re.match(r'(outreg2|esttab)\b',s):continue
        if s.startswith('save "$OUT/'):continue
        if s.startswith('global OUT '):cmd=f'global OUT "{OUT.as_posix()}"'
        if s.startswith('log using '):cmd=f'log using "{OUT.as_posix()}/run.log", text replace'
        out.append(cmd)
        if s.startswith('use "$ROOT/'):
            out.append(helper)
            # Same underlying measure, not a change of definition.
            out.append('clonevar gunning_fog_win3 = fog_win3')
        if s.startswith('pwcorr '):out.append(stats)
        if re.match(r'^(reghdfe|ivreghdfe|regress|stcox|streg|ppmlhdfe|margins)\s',s):
            out.append('word_capture')
            meta.append(dict(line=start,section=section,command=cmd))
    out+=['postclose wordcapture','display "V3_WORD_COMPLETE"','log close']
    (OUT/'execute.do').write_text('\n'.join(out)+'\n',encoding='utf-8')
    (OUT/'commands.json').write_text(json.dumps(meta,indent=2))
    (OUT/'provenance.json').write_text(json.dumps({'source':str(SRC),'source_sha256':hashlib.sha256(SRC.read_bytes()).hexdigest(),'doc':str(DOC),'doc_sha256':hashlib.sha256(DOC.read_bytes()).hexdigest(),'runtime_adjustments':['omit Excel exports and intermediate analysis-data save','alias gunning_fog_win3 to existing fog_win3; no outcome recoding']},indent=2))
    print('Prepared',len(meta),'static commands (loops expand at runtime)')

def build():
    import pandas as pd
    from lxml import etree as ET
    from copy import deepcopy
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    C=pd.read_stata(OUT/'coefficients.dta',convert_categoricals=False)
    assert C.model.nunique()==81
    assert 'V3_WORD_COMPLETE' in (OUT/'run.log').read_text(errors='replace')
    provenance=json.loads((OUT/'provenance.json').read_text())
    assert hashlib.sha256(SRC.read_bytes()).hexdigest()==provenance['source_sha256']
    assert hashlib.sha256(DOC.read_bytes()).hexdigest()==provenance['doc_sha256']
    D=Document(DOC)
    before=[ET.tostring(t._tbl) for t in D.tables]
    audit=[]; changed=set()
    def text(el):return ''.join(el.xpath('.//w:t/text()'))
    def rewrite(el,value):
        nodes=el.xpath('.//w:t')
        if not nodes:
            if not value:return
            runs=el.xpath('.//w:r')
            if runs:r=runs[0]
            else:
                p=el.xpath('.//w:p')[0];r=OxmlElement('w:r');p.append(r)
            n=OxmlElement('w:t');r.append(n);nodes=[n]
        # Preserve every existing run and its formatting; distribute replacement
        # over the original text slots to avoid flattening italic/bold labels.
        pos=0
        for k,n in enumerate(nodes):
            length=len(n.text or '')
            n.text=value[pos:] if k==len(nodes)-1 else value[pos:pos+length]
            pos+=length
            if n.text and (n.text[0].isspace() or n.text[-1].isspace()):n.set(qn('xml:space'),'preserve')
    def cell(t,r,c,v):
        el=D.tables[t].rows[r]._tr.tc_lst[c]
        old=text(el);v=str(v)
        if old!=v:
            rewrite(el,v);changed.add(t)
            audit.append(dict(table=t,row=r,col=c,old=old,new=v))
    def num(x,integer=False):
        if pd.isna(x):return '—'
        if integer:return f'{int(round(x)):,}'
        if abs(x)<.0005:x=0
        return f'{x:.3f}'
    def get(mid,term):
        z=C[C.model.eq(mid)&C.se.gt(0)]
        if term.startswith('margin:'):
            stage=term.split(':')[1]
            z=z[z.term.str.match(r'1\.VD:'+stage+r'(?:bn)?\.fundstage$')]
        else:
            terms=[term,'1.'+term]
            if 'priorpatentdummy_VD' in term:
                terms += [s.replace('priorpatentdummy_VD','patentdummybyfiling') for s in terms]
            z=z[z.term.str.split(':').str[-1].isin(terms)]
        assert len(z)<=1,(mid,term,z.to_dict())
        return None if z.empty else z.iloc[0]
    def put(t,r,c,mid,term):
        z=get(mid,term)
        if z is None:b=p=''
        else:
            b=num(z.b)+('***' if z.p<.01 else '**' if z.p<.05 else '*' if z.p<.1 else '')
            p='('+num(z.p)+')'
        cell(t,r,c,b);cell(t,r+1,c,p)
    def stats(t,r,c,mid,field):
        z=C[C.model.eq(mid)].iloc[0]
        cell(t,r,c,num(z[field],field in ['N','fail']))
    common={2:'VD',4:'1.VD#1.priorpatentdummy_VD',6:'1.VD#1.latestage_VD',8:'priorpatentdummy_VD',10:'latestage_VD',12:'ln_MoneyRaised_VD',14:'NumofFundingRounds_VD',16:'_cons'}
    def regression(t,mids,terms=common,nrow=22,rrow=23):
        for col,mid in enumerate(mids,1):
            for row,term in terms.items():put(t,row,col,mid,term)
            stats(t,nrow,col,mid,'N')
            if rrow is not None:
                ispp=C[C.model.eq(mid)].iloc[0]['cmd']=='ppmlhdfe'
                stats(t,rrow,col,mid,'r2p' if ispp else 'r2')
                if ispp:cell(t,rrow,0,'Pseudo R-squared')

    # 1. Summary statistics; funding is displayed in USD millions in template.
    S=pd.read_stata(OUT/'summary.dta').set_index('variable')
    variables=list(S.index)
    for row,v in enumerate(variables,1):
        for col,field in enumerate(['N','mean','sd','min','max'],1):
            value=S.loc[v,field]
            if v=='MoneyRaised_VD' and field!='N':value/=1e6
            cell(0,row,col,num(value,field=='N'))
    # 2. Pairwise correlation: source uses funding LEVEL, not log funding.
    R=pd.read_stata(OUT/'correlations.dta').set_index(['variable1','variable2'])
    for i,v in enumerate(variables):
        for j,w in enumerate(variables[:6]):
            if j<=i:
                z=R.loc[(v,w)];cell(1,i+1,j+2,num(z.r)+('*' if i!=j and z.p<.05 else ''))
    for i in range(6,12):
        for j in range(6,i+1):
            z=R.loc[(variables[i],variables[j])]
            cell(2,i-5,j-4,num(z.r)+('*' if i!=j and z.p<.05 else ''))
    cell(1,5,1,'Total Funding at VD')
    regression(3,range(1,5))
    stage={2:'VD',4:'1.VD#1.priorpatentdummy_VD',6:'1.VD#2.fundstage',8:'1.VD#3.fundstage',10:'1.VD#4.fundstage',12:'priorpatentdummy_VD',14:'2.fundstage',16:'3.fundstage',18:'4.fundstage',20:'ln_MoneyRaised_VD',22:'NumofFundingRounds_VD',24:'_cons'}
    regression(4,[5,7,9,11],stage,29,30)
    for col,mid in enumerate([6,8,10,12],1):
        for stageid,row in enumerate([2,4,6,8],1):put(5,row,col,mid,f'margin:{stageid}')
    regression(6,range(13,17));regression(7,range(17,21))
    short={3:'VD',5:'priorpatentdummy_VD',7:'latestage_VD',9:'ln_MoneyRaised_VD',11:'NumofFundingRounds_VD',13:'_cons'}
    regression(8,range(21,24),short,19,20)
    for col,mid in enumerate(range(24,30),1):
        for row,term in list(short.items())[:-1]:put(9,row,col,mid,term)
        stats(9,17,col,mid,'fail');stats(9,18,col,mid,'N')
    # IV: model 30 is second stage; model 31 is first stage.
    for col,mid in enumerate([30,31],1):
        for row,term in {7:'priorpatentdummy_VD',9:'latestage_VD',11:'ln_MoneyRaised_VD',13:'NumofFundingRounds_VD'}.items():put(16,row,col,mid,term)
        stats(16,19,col,mid,'N');stats(16,20,col,mid,'F');stats(16,21,col,mid,'r2')
    put(16,3,2,31,'ln_vdsupply');put(16,5,1,30,'VD')
    log=(OUT/'run.log').read_text(errors='replace')
    cd=re.search(r'Weak identification test \(Cragg-Donald Wald F statistic\):\s*([\d.]+)',log)
    assert cd;cell(16,22,1,num(float(cd.group(1))))
    regression(17,range(32,36))
    finance={2:'VD',4:'priorpatentdummy_VD',6:'latestage_VD',8:'ln_MoneyRaised_VD',10:'NumofFundingRounds_VD',12:'_cons'}
    regression(18,range(36,39),finance,18,19)
    regression(19,range(39,42),finance,18,19)
    # Do not insert non-existent ordinary R2 values for survival estimators.
    cell(19,19,1,'');cell(19,19,2,'')
    regression(20,range(42,45),finance,18,19)
    regression(21,range(45,50),common,23,24)
    for ti,start in zip(range(24,32),range(50,82,4)):regression(ti,range(start,start+4))
    # Only update labels in recalculated tables; excluded tables stay byte-identical.
    for ti in sorted(changed):
        for row in D.tables[ti].rows:
            for tc in row._tr.tc_lst:
                old=text(tc)
                new=old.replace('Early Stage','Late Stage').replace('No Prior Patent','Prior Patent')
                if new!=old:rewrite(tc,new)
    # Relevant notes only: no global replacement (would alter excluded tables/figures).
    current=''
    note_changes=[]
    for p in D.paragraphs:
        match=re.match(r'TABLE\s+([A]?\d+(?:-\d+)?)\.',p.text,re.I)
        if match:current=match.group(1)
        if not p.text.startswith('Note'):continue
        if current in ['3','A4']:
            old=p.text;new=old.replace('Early Stage','Late Stage')
            if new!=old:rewrite(p._p,new);note_changes.append(current)
        if current=='A2':
            old=text(p._p);new=re.sub(r'N\s*=\s*758', 'N=710',old)
            if old!=new:rewrite(p._p,new);note_changes.append(current)
    excluded=[10,11,12,13,14,15,22,23]
    for ti in excluded:assert ET.tostring(D.tables[ti]._tbl)==before[ti],('excluded modified',ti)
    for ti in [32,33]:assert ET.tostring(D.tables[ti]._tbl)==before[ti]
    xml=ET.tostring(D._element,xml_declaration=True,encoding='UTF-8',standalone=True)
    with zipfile.ZipFile(DOC) as src,zipfile.ZipFile(FINAL,'w',zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():dst.writestr(item,xml if item.filename=='word/document.xml' else src.read(item.filename))
    with zipfile.ZipFile(DOC) as a,zipfile.ZipFile(FINAL) as b:
        assert a.namelist()==b.namelist()
        assert all(a.read(n)==b.read(n) for n in a.namelist() if n!='word/document.xml')
    # Verify all recorded replacements by reading the saved document again.
    F=Document(FINAL)
    latest={(c['table'],c['row'],c['col']):c for c in audit}
    for change in latest.values():
        value=text(F.tables[change['table']].rows[change['row']]._tr.tc_lst[change['col']])
        assert value==change['new'].replace('Early Stage','Late Stage').replace('No Prior Patent','Prior Patent'),change
    for tag in ['w:tblPr','w:tblGrid','w:trPr','w:tcPr','w:pPr','w:rPr','w:sectPr']:
        a=D._element.xpath('.//'+tag); b=Document(DOC)._element.xpath('.//'+tag)
        assert [ET.tostring(x) for x in a]==[ET.tostring(x) for x in b],('formatting changed',tag)
    report={'models':81,'numeric_cells_updated':len(latest),'changed_table_indices':sorted(changed),'excluded_tables_identical':True,'formatting_properties_identical':True,'all_non_document_xml_parts_identical':True,'note_changes':note_changes,'final':str(FINAL),'changes':list(latest.values())}
    (OUT/'document_audit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print('Saved',FINAL,'cells',len(audit))

if __name__=='__main__':
    if sys.argv[1]=='prepare':prepare()
    elif sys.argv[1]=='build':build()
