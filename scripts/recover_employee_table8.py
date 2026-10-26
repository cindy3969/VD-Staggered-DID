from pathlib import Path
import pandas as pd
import numpy as np
from openpyxl import load_workbook
B=Path(__file__).resolve().parents[1]
O=B/'outputs/psm_final_2820_package/postvd_size_research'
d=pd.read_stata(B/'260122 winsorization_stage4.dta',convert_categoricals=False)
m=pd.read_stata(B/'outputs/psm_final_2820_package/table3_current_2820_cross_section_v2.dta',convert_categoricals=False)
q=d[d.lastdeal2vd1.between(1,3)&d.private.eq(1)].copy()
X=np.column_stack([np.ones(len(q)),q.VD,pd.get_dummies(q.lastdeal2vd1,drop_first=True).astype(float)])
y=q.emp2025_2.to_numpy();b=np.linalg.lstsq(X,y,rcond=None)[0];e=y-X@b
print('MODEL1 coefficients',b,'N',len(y),'R2',1-e@e/((y-y.mean())@(y-y.mean())))
print('OLD 611 firms matching current cohort',q.OrganizationName.isin(m.OrganizationName).sum())
q.to_csv(O/'recovered_table8_original_611.csv',index=False)
for f in ['employees by VD group.xlsx','cross sectional 2820.xlsx','exact matched sample v7 unmatched 4.xlsx']:
    w=load_workbook(Path('C:/Users/user/iCloudDrive/5학기/연구미팅/20240108-2010 to 2015 firms')/f,read_only=True,data_only=True)
    print(f,[(s.title,s.max_row,s.max_column) for s in w.worksheets])
    for r in w.active.iter_rows(min_row=1,max_row=2,values_only=True): print(r)
