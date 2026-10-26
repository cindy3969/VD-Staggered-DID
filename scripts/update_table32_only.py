from pathlib import Path
import sys,re,json,zipfile,shutil
from io import BytesIO
from lxml import etree
import pandas as pd
from docx import Document
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
BASE=ROOT/'outputs/table_run_v2_20261020'
OUT=BASE/'table32_revision'
OUT.mkdir(exist_ok=True)
DOC=BASE/'VD Table 20261020 original 2820 v2.docx'
SRC=ROOT/'outputs/psm_final_2820_package/RUN_TABLE3_PATENT_APPLICATION_WIN3_v2.do'

if sys.argv[1]=='prepare':
    prior=(BASE/'execute_no_excel.do').read_text(encoding='utf-8')
    # Reuse verified setup and capture routine, without re-estimating other tables.
    setup=prior.split('* Table 3: Patent quantity')[0]
    setup=setup.replace(f'global OUT "{BASE.as_posix()}"',f'global OUT "{OUT.as_posix()}"')
    source=SRC.read_text(encoding='utf-8-sig')
    merge=source.split('* Table 3-2:')[1].split('* Panel B:')[0]
    merge='* Table 3-2:'+merge
    merge='\n'.join(x for x in merge.splitlines() if not x.startswith('save "$OUT'))
    block=source.split('* Panel B:')[1].split('* Table 4:')[0]
    block='* Panel B:'+block
    lines=block.splitlines(); commands=[]; i=0; m=0; output=[]
    while i<len(lines):
        group=[lines[i]];i+=1
        while group[-1].rstrip().endswith('///') and i<len(lines):
            group.append(lines[i]);i+=1
        cmd='\n'.join(group)
        if cmd.startswith('reghdfe '):m+=1
        if m<=2 and (cmd.startswith('testparm ') or cmd.startswith('lincom ')):continue
        output.append(cmd)
        if cmd.startswith('reghdfe '):
            output.append(f'capture_model, id({2*m-1})')
            commands.append(cmd)
        if cmd.startswith('margins '):output.append(f'capture_model, id({2*m})')
    assert m==4
    (OUT/'run_table32.do').write_text(setup+'\n'+merge+'\n'+'\n'.join(output)+
          '\npostclose capture_results\ndisplay "TABLE32_COMPLETE"\nlog close\n',encoding='utf-8')
    (OUT/'specifications.json').write_text(json.dumps(commands,indent=2))
    print(OUT/'run_table32.do')
    sys.exit()

C=pd.read_stata(OUT/'coefficients.dta')
assert C.model.nunique()==8
assert 'TABLE32_COMPLETE' in (OUT/'RUN_TABLE3_FULL_v2.log').read_text(errors='replace')
backup=OUT/'before_table32_update.docx'
if not backup.exists():shutil.copy2(DOC,backup)
D=Document(DOC)
body=D._element.body
original=[etree.tostring(e) for e in body]
changed=set(); audit=[]
tables=[]
for t in D.tables:
    prev=t._tbl.getprevious()
    title=''.join(prev.itertext()) if prev is not None else ''
    if 'Table 3-2.' in title:tables.append(t)
assert len(tables)==3

def text(el,value):
    ts=el.findall('.//'+qn('w:t'))
    assert ts
    ts[0].text=value
    for x in ts[1:]:x.text=''

def cell(t,r,c,value):
    el=t.cell(r,c)._tc
    ts=el.findall('.//'+qn('w:t'))
    if not ts:
        from docx.oxml import OxmlElement
        p=t.cell(r,c).paragraphs[0]._p
        run=OxmlElement('w:r');tx=OxmlElement('w:t');run.append(tx);p.append(run)
    text(el,value)

def find(mid,term):
    z=C[C.model.eq(mid)]
    z=z[z.term.str.split(':').str[-1].isin([term,'1.'+term]) & z.se.gt(0)]
    assert len(z)<=1
    return None if z.empty else z.iloc[0]

def put(t,r,c,mid,term):
    z=find(mid,term);b=p=''
    if z is not None:
        b=f'{0 if abs(z.b)<.0005 else z.b:.3f}'
        b+='***' if z.p<.01 else '**' if z.p<.05 else '*' if z.p<.1 else ''
        p='(<0.001)' if z.p<.001 else f'({z.p:.3f})'
    cell(t,r,c,b);cell(t,r+1,c,p)
    audit.append(dict(model=mid,term=term,b=b,p=p))

mapping={'VD':'VD','VD × Series A':'1.VD#2.fundstage','VD × Series B':'1.VD#3.fundstage',
 'VD × Series C+':'1.VD#4.fundstage','VD × Prior Patent':'1.VD#1.priorpatentdummy_VD',
 'Series A':'2.fundstage','Series B':'3.fundstage','Series C+':'4.fundstage',
 'Prior Patent':'priorpatentdummy_VD','Log(Total Funding at VD)':'ln_MoneyRaised_VD',
 'Funding Rounds at VD':'NumofFundingRounds_VD','Constant':'_cons'}
common=('Columns (1)-(4): baseline; VD x Prior Patent; VD x funding stage; both interactions. '
        'Seed/pre-Series is the stage reference category. Prior Patent=1 denotes prior patent ownership. '
        'All models include region, industry and founding-year fixed effects and robust standard errors. '
        'Two-sided p-values in parentheses. * p<0.10, ** p<0.05, *** p<0.01.')
for k,t in enumerate(tables):
    changed.add(body.index(t._tbl))
    note=t._tbl.getnext();assert note.tag==qn('w:p')
    changed.add(body.index(note))
    if k<2:
        remove=[]
        for r in range(2,len(t.rows),2):
            label=t.cell(r,0).text
            if label in ['Early Stage','VD × Early Stage']:
                remove.extend([t.rows[r]._tr,t.rows[r+1]._tr]);continue
            assert label in mapping,label
            for col,mid in enumerate([1,3,5,7],1):put(t,r,col,mid,mapping[label])
        for row in remove:t._tbl.remove(row)
        n=[int(C[C.model.eq(mid)].N.iloc[0]) for mid in [1,3,5,7]]
        assert len(set(n))==1
        text(note,common+f' N={n[0]:,} in all columns. Stage assignments include source imputations.')
    else:
        for r,stage in [(2,1),(4,2),(6,3),(8,4)]:
            term=('1bn' if stage==1 else str(stage))+'.fundstage'
            for col,mid in enumerate([2,4,6,8],1):put(t,r,col,mid,term)
        text(note,'Average discrete VD effects from margins fundstage, dydx(1.VD). '
             'Columns (1)-(2) have no VD x funding-stage interaction, so the average VD effect is identical across stages. '+common)
assert len(original)==len(body)
assert all(original[i]==etree.tostring(e) for i,e in enumerate(body) if i not in changed)
buf=BytesIO();D.save(buf)
with zipfile.ZipFile(buf) as z:xml=z.read('word/document.xml')
temp=OUT/'updated.docx'
with zipfile.ZipFile(DOC) as src,zipfile.ZipFile(temp,'w',zipfile.ZIP_DEFLATED) as dst:
    for entry in src.infolist():
        dst.writestr(entry,xml if entry.filename=='word/document.xml' else src.read(entry.filename))
with zipfile.ZipFile(DOC) as src,zipfile.ZipFile(temp) as dst:
    assert all(src.read(n)==dst.read(n) for n in src.namelist() if n!='word/document.xml')
try:
    temp.replace(DOC)
    final=DOC
except PermissionError:
    final=BASE/'VD Table 20261020 original 2820 v2 Table32 updated.docx'
    temp.replace(final)
(OUT/'audit.json').write_text(json.dumps({'changed_body_elements':sorted(changed),'cells':audit},indent=2))
print('Updated only three Table 3-2 panels and their notes; all other body elements and package parts unchanged.')
print(final)
