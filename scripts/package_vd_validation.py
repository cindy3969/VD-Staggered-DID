from pathlib import Path
import json,zipfile,datetime
import openpyxl
R=Path('C:/Users/user/Documents/VD'); O=R/'outputs/vd_lender_validation_2820'
wb=openpyxl.load_workbook(O/'VD_lender_validation_2820.xlsx',read_only=True,data_only=True)
assert set(wb.sheetnames)=={'Summary','Firms','Deals','Providers','Public checks','Missingness','Methods'}
for name,n in [('Firms',2821),('Deals',1899),('Providers',711),('Public checks',9)]:
    assert sum(row[0] is not None for row in wb[name].iter_rows(values_only=True))==n
assert wb['Summary']['C9'].value==715
assert abs(wb['Summary']['E9'].value-715/1410)<1e-10
assert isinstance(wb['Deals']['G2'].value,datetime.datetime),repr(wb['Deals']['G2'].value)
assert wb['Deals']['G2'].value.strftime('%Y-%m-%d')=='2022-05-12'
for sheet in wb:
    for row in sheet:
        for c in row:
            assert c.data_type!='e',(sheet.title,c.coordinate,c.value)
wb.close()
v=json.loads((O/'verification.json').read_text())
v['xlsx_reopened_and_checked']=True
v['excel_error_cells']=0
v['excel_dates_typed']=True
(O/'verification.json').write_text(json.dumps(v,indent=2),encoding='utf-8')
selected=[x for x in O.iterdir() if x.suffix in ['.csv','.dta','.do','.html','.xlsx'] and not x.name.startswith('preview')]
selected += [O/'manifest.json',O/'verification.json']
scripts=['validate_vd_lenders.py','vd_external_deal_evidence.py','report_vd_validation.py','build_vd_validation_workbook.mjs','package_vd_validation.py']
with zipfile.ZipFile(O/'VD_validation_complete_package.zip','w',zipfile.ZIP_DEFLATED) as z:
    for x in selected:z.write(x,x.name)
    for x in scripts:z.write(R/'scripts'/x,'reproduction/'+x)
print(json.dumps(v));print('Packaged',len(selected),'outputs plus',len(scripts),'scripts')
