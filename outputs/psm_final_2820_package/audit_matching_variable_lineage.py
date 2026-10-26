import re
import pandas as pd
import numpy as np
from openpyxl import load_workbook

ROOT = r"C:\Users\user\Documents\VD\outputs\psm_final_2820_package"
V4 = r"C:\Users\user\iCloudDrive\5학기\연구미팅\20240108-2010 to 2015 firms\2010 2015 Crunchbase deals merged_v4.xlsx"
V6 = r"C:\Users\user\iCloudDrive\5학기\연구미팅\20240108-2010 to 2015 firms\exact matched sample v6 unmatched.xlsx"
FINAL = r"C:\Users\user\iCloudDrive\5학기\연구미팅\20240108-2010 to 2015 firms\251111 2820 cross sectional.dta"

def key(x):
    return re.sub(r"[^a-z0-9]+", "", "" if pd.isna(x) else str(x).casefold())

v6 = pd.read_excel(V6)
v6["key"] = v6["Organization Name"].map(key)
f = pd.read_stata(FINAL, convert_categoricals=False)
f["key"] = f["OrganizationName"].map(key)

print("FINAL columns likely matching variables:")
for c in f.columns:
    if any(x in c.lower() for x in ["fund", "grant", "patent", "found", "region", "industry", "timetofirst", "round"]):
        print(c)

print("\nV6-vs-final matches by common variable candidates")
merged = f.merge(v6, on="key", how="left", suffixes=("_final", "_v6"))
for a,b in [
    ("Foundingyear", "Founding_Year"),
    ("TimeToFirst_VC_final", "TimeToFirst_VC_v6"),
    ("MoneyRaised_VD", "Total_Funding_Amount_FirstVD"),
    ("NumofFundingRounds_VD", "NumberOfFunding_Rounds_FirstVD"),
]:
    if a in merged and b in merged:
        x=pd.to_numeric(merged[a],errors="coerce"); y=pd.to_numeric(merged[b],errors="coerce")
        ok=x.notna()&y.notna(); print(a,b,"n",ok.sum(),"exact",np.isclose(x[ok],y[ok],rtol=1e-10,atol=1e-8).sum(),"maxdiff",float((x[ok]-y[ok]).abs().max()))

# Read relevant V4 deal-level fields and show the first matched treated/control histories.
wb=load_workbook(V4,read_only=True,data_only=True); ws=wb.active
rows=[]
for r in ws.iter_rows(min_row=3,values_only=True):
    rows.append({
        "name":r[3],"key":key(r[3]),"weight":max(int(r[2]==1),int(r[36]==1)),
        "found_year":r[6],"deal_date":r[7],"deal_year":r[8],"funding_type":r[10],
        "first_vc":r[13],"time_first_vc":r[14],"debt":r[15],"vd_backing":r[16],"first_vd":r[17],
        "time_first_vd":r[30],"vd_year_col32":r[31],"vd_year_deal":r[33],
        "rounds_firstvd":r[34],"funding_firstvd":r[35],"vd_year_col44":r[43],
        "amount":r[44],"region":r[74],"industry_group":r[85],"industry_class":r[86]
    })
deals=pd.DataFrame(rows)

firm=deals.groupby("key",as_index=False).agg(
    name=("name","first"),weight=("weight","max"),vd_backing=("vd_backing","max"),
    n_deals=("deal_year","size"),found_year=("found_year","first"),
    first_vc_year=("deal_year",lambda x: np.nan),
    vd_year32=("vd_year_col32",lambda x: next((v for v in x if pd.notna(v) and v!='`'),np.nan)),
    vd_year44=("vd_year_col44",lambda x: next((v for v in x if pd.notna(v) and v!='`'),np.nan)),
    rounds_firstvd=("rounds_firstvd",lambda x: next((v for v in x if pd.notna(v) and v!=''),np.nan)),
    funding_firstvd=("funding_firstvd",lambda x: next((v for v in x if pd.notna(v) and v!=''),np.nan)),
    time_first_vc=("time_first_vc",lambda x: next((v for v in x if pd.notna(v) and v!=''),np.nan)),
    region=("region","first"),industry_group=("industry_group","first"),industry_class=("industry_class","first")
)

firm=firm.merge(v6[["key","VD Year","Founding_Year","VD_Backing","TimeToFirst_VC","Grant_VD","Total_Funding_Amount_FirstVD","NumberOfFunding_Rounds_FirstVD","Region2","Industry2"]],on="key",how="outer")
firm.to_csv(ROOT+r"\matching_variable_lineage_all_firms.csv",index=False,encoding="utf-8-sig")

print("\nV4 firm counts",len(firm),"weight",int(firm.weight.fillna(0).sum()))
print("V4 nonmissing VD year col32 by treatment")
print(firm.groupby("VD_Backing")["vd_year32"].apply(lambda x:x.notna().sum()))
print("V6 VD Year distribution by treatment")
print(firm.groupby("VD_Backing")["VD Year"].describe()[["count","mean","min","max"]])

sample_keys=[]
for vd in [1,0]:
    q=firm[(firm.weight.eq(1))&(firm.VD_Backing.eq(vd))].head(2)
    sample_keys += q.key.tolist()
deals[deals.key.isin(sample_keys)].sort_values(["key","deal_year"]).to_csv(ROOT+r"\matching_variable_deal_history_examples.csv",index=False,encoding="utf-8-sig")
