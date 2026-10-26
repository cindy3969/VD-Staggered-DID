from pathlib import Path
from openpyxl import load_workbook

root = Path(r"C:\Users\user\iCloudDrive\5학기\연구미팅\20240108-2010 to 2015 firms")
names = [
    "exact matched sample.xlsx",
    "exact matched sample v9 unmatched.xlsx",
    "exact matched sample v7 unmatched_recovered.xlsx",
    "exact matched sample v7 unmatched.xlsx",
    "exact matched sample v7 unmatched 4.xlsx",
    "exact matched sample v7 unmatched 3.xlsx",
    "exact matched sample v7 unmatched 2.xlsx",
    "exact matched sample v6 unmatched.xlsx",
    "exact matched sample v5 unmatched.xlsx",
    "exact matched sample v3.xlsx",
    "cross sectional 2820.xlsx",
]
needles = ("pscore", "score", "weight", "_id", "_n1", "_nn", "pdif", "pair", "match")
for name in names:
    p = root / name
    print("\n" + "=" * 100)
    print(f"FILE: {name} | exists={p.exists()} | size={p.stat().st_size if p.exists() else None}")
    if not p.exists():
        continue
    try:
        wb = load_workbook(p, read_only=True, data_only=False)
    except Exception as e:
        print(f"OPEN ERROR: {e}")
        continue
    for ws in wb.worksheets:
        header = []
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 10), values_only=True):
            header.extend([str(x).strip() for x in row if x is not None])
        hits = sorted({h for h in header if any(n in h.lower() for n in needles)})
        print(f"  SHEET: {ws.title!r} rows={ws.max_row} cols={ws.max_column}")
        print(f"    first-row headers: {[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))][:30]}")
        print(f"    PSM-like header hits in first 10 rows: {hits[:50]}")
    wb.close()
