from openpyxl import load_workbook
from collections import Counter

f = r"C:\Users\user\iCloudDrive\5학기\연구미팅\20240108-2010 to 2015 firms\2010 2015 Crunchbase deals merged_v4.xlsx"
s = load_workbook(f, read_only=True, data_only=True).active
firms = {}
for row in s.iter_rows(min_row=3, values_only=True):
    name, url, w3, w37, vd = row[3], row[1], row[2], row[36], row[16]
    if name:
        key = str(url).strip().lower() if url else str(name).strip().lower()
        rec = firms.setdefault(key, {"name": name, "w3": set(), "w37": set(), "vd": set()})
        rec["w3"].add(w3)
        rec["w37"].add(w37)
        rec["vd"].add(vd)
print("unique firms", len(firms))
for col in ["w3", "w37"]:
    sel = [r for r in firms.values() if 1 in r[col]]
    vd_counts = Counter(v for r in sel for v in r["vd"] if v is not None)
    print(col, "matched unique", len(sel), "VD", dict(vd_counts))
print("w3-w37 discord", sum((1 in r["w3"]) != (1 in r["w37"]) for r in firms.values()))

# Alternative counts by organization name only and diagnostics for inconsistent treatment.
by_name = {}
for r in firms.values():
    key = str(r["name"]).strip().casefold()
    z = by_name.setdefault(key, {"name": r["name"], "w3": set(), "vd": set()})
    z["w3"].update(r["w3"])
    z["vd"].update(r["vd"])
sel = [r for r in by_name.values() if 1 in r["w3"]]
print("unique organization names", len(by_name))
print("matched by name", len(sel), "VD", dict(Counter(v for r in sel for v in r["vd"] if v is not None)))
print("matched inconsistent VD", [(r["name"], sorted(str(v) for v in r["vd"])) for r in sel if len({v for v in r["vd"] if v is not None}) > 1][:30])
