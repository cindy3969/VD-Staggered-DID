from pathlib import Path
from openpyxl import load_workbook
from collections import Counter

p=Path(r"C:\Users\user\iCloudDrive\5학기\연구미팅\20240108-2010 to 2015 firms\exact matched sample v6 unmatched.xlsx")
wb=load_workbook(p,read_only=True,data_only=True)
ws=wb["Sheet1"]
headers=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]
for i,h in enumerate(headers,1): print(i,repr(h))
lookup={str(h).strip().lower():i for i,h in enumerate(headers) if h is not None}
fields=["VD_Backing","Founding_Year","VD Year","TimeToFirst_VC","NumberOfFunding_Rounds_FirstVD","NumberOfFunding_Rounds_FirstVC","Total_Funding_Amount_FirstVD","Total_Funding_Amount_FirstVC","Region2","Industry2","Grant_VD"]
idx={f:lookup.get(f.lower()) for f in fields}
print("indices",idx)
counts=Counter(); missing=Counter(); rows=0
for row in ws.iter_rows(min_row=2,values_only=True):
    if not any(v is not None for v in row): continue
    rows+=1
    vi=idx["VD_Backing"]
    counts[row[vi] if vi is not None else None]+=1
    for f,i in idx.items():
        if i is None or i>=len(row) or row[i] in (None,""): missing[f]+=1
print("nonblank rows",rows)
print("VD counts",counts)
print("missing",missing)
wb.close()
