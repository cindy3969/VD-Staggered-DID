import re
import pandas as pd
from openpyxl import load_workbook

def key(x):
    return re.sub(r"[^a-z0-9]+", "", str(x).casefold())

x = r"C:\Users\user\iCloudDrive\5학기\연구미팅\20240108-2010 to 2015 firms\2010 2015 Crunchbase deals merged_v4.xlsx"
dta = r"C:\Users\user\iCloudDrive\5학기\연구미팅\20240108-2010 to 2015 firms\251111 2820 cross sectional.dta"
s = load_workbook(x, read_only=True, data_only=True).active
weighted = {}
for r in s.iter_rows(min_row=3, values_only=True):
    if r[3] is not None and (r[2] == 1 or r[36] == 1):
        weighted[key(r[3])] = r[3]
f = pd.read_stata(dta, convert_categoricals=False)
final = {key(v): v for v in f.OrganizationName}
print("weight names", len(weighted), "final names", len(final), "overlap", len(set(weighted)&set(final)))
print("weight only", sorted(weighted[k] for k in set(weighted)-set(final)))
print("final only", sorted(final[k] for k in set(final)-set(weighted)))
