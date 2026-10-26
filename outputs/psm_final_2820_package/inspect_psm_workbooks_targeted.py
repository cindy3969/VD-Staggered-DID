from pathlib import Path
from openpyxl import load_workbook
from collections import Counter

root=Path(r"C:\Users\user\iCloudDrive\5학기\연구미팅\20240108-2010 to 2015 firms")
files=[
 "Cross sectional 2010 2015 Crunchbase deals merged_v2.xlsx",
 "exact matched sample v5 unmatched.xlsx",
 "exact matched sample v6 unmatched.xlsx",
 "exact matched sample.xlsx",
]
for name in files:
    print("\n"+"="*90+"\nFILE",name)
    wb=load_workbook(root/name,read_only=True,data_only=True)
    for ws in wb.worksheets:
        headers=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]
        hits=[(i+1,h) for i,h in enumerate(headers) if h is not None and any(k in str(h).lower() for k in ["pscore","weight","_id","_n1","_nn","pdif","matched"])]
        print("SHEET",ws.title,"rows",ws.max_row,"cols",ws.max_column,"hits",hits)
        if ws.max_row<=20000 and hits:
            lookup={str(h).strip().lower():i for i,h in enumerate(headers) if h is not None}
            vd_i=lookup.get("vd_backing",lookup.get("vd"))
            w_i=lookup.get("_weight")
            p_i=lookup.get("_pscore",lookup.get("pscore"))
            counts=Counter(); wc=Counter(); pn=0; pmin=None; pmax=None
            for row in ws.iter_rows(min_row=2,values_only=True):
                vd=row[vd_i] if vd_i is not None and vd_i<len(row) else None
                w=row[w_i] if w_i is not None and w_i<len(row) else None
                p=row[p_i] if p_i is not None and p_i<len(row) else None
                counts[vd]+=1; wc[(vd,w)]+=1
                if isinstance(p,(int,float)):
                    pn+=1; pmin=p if pmin is None else min(pmin,p); pmax=p if pmax is None else max(pmax,p)
            print("  VD",counts)
            print("  VD_WEIGHT",wc)
            print("  PSCORE",pn,pmin,pmax)
    wb.close()
