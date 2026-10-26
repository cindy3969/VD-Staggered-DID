from pathlib import Path
from openpyxl import load_workbook
from collections import Counter

p=Path(r"C:\Users\user\iCloudDrive\5학기\연구미팅\20240108-2010 to 2015 firms\exact matched sample.xlsx")
wb=load_workbook(p,read_only=True,data_only=True)
ws=wb[wb.sheetnames[0]]
headers=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]
print("headers")
for i,h in enumerate(headers,1): print(i,repr(h))
cols={str(h):i for i,h in enumerate(headers) if h is not None}
for key in ["VD_Backing","_pscore","_weight","_id","_n1","_nn","_pdif","pscore","pscore2","OrganizationName"]:
    print("column",key,cols.get(key))
vd_i=cols.get("VD_Backing")
w_i=cols.get("_weight")
counts=Counter(); wc=Counter()
for row in ws.iter_rows(min_row=2,values_only=True):
    vd=row[vd_i] if vd_i is not None and vd_i<len(row) else None
    w=row[w_i] if w_i is not None and w_i<len(row) else None
    counts[vd]+=1
    wc[(vd,w)]+=1
print("VD counts",counts)
print("VD-weight counts",wc)
wb.close()
