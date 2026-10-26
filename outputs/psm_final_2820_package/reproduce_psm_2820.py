import re
import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = r"C:\Users\user\Documents\VD\outputs\psm_final_2820_package"
PRE = r"C:\Users\user\iCloudDrive\5학기\연구미팅\20240108-2010 to 2015 firms\exact matched sample v6 unmatched.xlsx"
DEALS = r"C:\Users\user\iCloudDrive\5학기\연구미팅\20240108-2010 to 2015 firms\2010 2015 Crunchbase deals merged_v4.xlsx"
FINAL = r"C:\Users\user\iCloudDrive\5학기\연구미팅\20240108-2010 to 2015 firms\251111 2820 cross sectional.dta"

def key(x):
    x = "" if pd.isna(x) else str(x)
    return re.sub(r"[^a-z0-9]+", "", x.casefold())

# Historical membership: weight==1, collapsed from deal to organization.
ws = load_workbook(DEALS, read_only=True, data_only=True).active
old = {}
for r in ws.iter_rows(min_row=3, values_only=True):
    if r[3] is None:
        continue
    k = key(r[3])
    z = old.setdefault(k, {"Organization Name": r[3], "old_weight": 0, "old_vd": 0})
    z["old_weight"] = max(z["old_weight"], int(r[2] == 1 or r[36] == 1))
    z["old_vd"] = max(z["old_vd"], int(r[16] == 1))
old = pd.DataFrame(old.values())
old["name_key"] = old["Organization Name"].map(key)

d = pd.read_excel(PRE)
d["name_key"] = d["Organization Name"].map(key)
# Restore the treatment definition used by the archived 2,820-firm analysis.
d.loc[d.name_key.isin([key("Dermira"), key("Eargo")]), "VD_Backing"] = 1
d.loc[d.name_key.isin([key("Cortex"), key("Nomi1")]), "VD_Backing"] = 0
d["ln_Total_Funding_Amount_FirstVD"] = np.log1p(d["Total_Funding_Amount_FirstVD"])

base = [
    "VD_Backing", "Organization Name", "Organization Name URL", "name_key",
    "Founding_Year", "Region2", "Industry2", "TimeToFirst_VC",
    "NumberOfFunding_Rounds_FirstVD", "Total_Funding_Amount_FirstVD",
    "ln_Total_Funding_Amount_FirstVD", "Grant_VD",
]
d = d[base].copy()
complete = [
    "Founding_Year", "Region2", "Industry2", "TimeToFirst_VC",
    "NumberOfFunding_Rounds_FirstVD", "ln_Total_Funding_Amount_FirstVD", "Grant_VD",
]
d["estimation_sample"] = d[complete].notna().all(axis=1).astype("int8")
e = d.loc[d.estimation_sample.eq(1)].copy()

# Equivalent to Stata factor-variable logit for the listed matching variables.
cat = pd.get_dummies(e[["Founding_Year", "Region2", "Industry2"]].astype(str), drop_first=True, dtype=float)
cont = e[["TimeToFirst_VC", "NumberOfFunding_Rounds_FirstVD", "ln_Total_Funding_Amount_FirstVD", "Grant_VD"]].astype(float)
X = pd.concat([cat, cont], axis=1)
X.insert(0, "const", 1.0)
y = e["VD_Backing"].astype(int)

# Binomial-logit IRLS, implemented in NumPy so the reproduction has no
# dependency on a separately installed statistical package.
xv = X.to_numpy(dtype=float)
yv = y.to_numpy(dtype=float)
b = np.zeros(xv.shape[1])
for _ in range(200):
    eta = np.clip(xv @ b, -30, 30)
    p = 1.0 / (1.0 + np.exp(-eta))
    w = np.clip(p * (1.0 - p), 1e-9, None)
    z = eta + (yv - p) / w
    sw = np.sqrt(w)
    b_new = np.linalg.pinv(xv * sw[:, None]) @ (z * sw)
    if np.max(np.abs(b_new - b)) < 1e-9:
        b = b_new
        break
    b = b_new
e["pscore_reestimated"] = 1.0 / (1.0 + np.exp(-np.clip(xv @ b, -30, 30)))

# Common support based on the intersection of treated and control score ranges.
t = e[e.VD_Backing.eq(1)]
c = e[e.VD_Backing.eq(0)]
lo = max(t.pscore_reestimated.min(), c.pscore_reestimated.min())
hi = min(t.pscore_reestimated.max(), c.pscore_reestimated.max())
e["common_support"] = e.pscore_reestimated.between(lo, hi).astype("int8")

def greedy(order):
    tt = e[(e.VD_Backing.eq(1)) & (e.common_support.eq(1))].copy()
    cc = e[(e.VD_Backing.eq(0)) & (e.common_support.eq(1))].copy()
    if order == "ascending":
        tt = tt.sort_values(["pscore_reestimated", "name_key"])
    elif order == "descending":
        tt = tt.sort_values(["pscore_reestimated", "name_key"], ascending=[False, True])
    else:
        tt = tt.sort_index()
    available = cc.index.to_numpy().copy()
    cp = cc.pscore_reestimated.to_numpy()
    rows = []
    for pair, (ti, tr) in enumerate(tt.iterrows(), 1):
        j = int(np.argmin(np.abs(cp - tr.pscore_reestimated)))
        ci = int(available[j])
        rows.append((pair, ti, ci, float(abs(cp[j] - tr.pscore_reestimated))))
        available = np.delete(available, j)
        cp = np.delete(cp, j)
    return rows

final = pd.read_stata(FINAL, convert_categoricals=False)
final["name_key"] = final["OrganizationName"].map(key)
old_treated = set(final.loc[final.VD.eq(1), "name_key"])
old_controls = set(final.loc[final.VD.eq(0), "name_key"])
summaries, candidates = [], {}
for order in ["original", "ascending", "descending"]:
    pairs = greedy(order)
    new_controls = set(e.loc[[x[2] for x in pairs], "name_key"])
    overlap = len(old_controls & new_controls)
    summaries.append({"treated_order": order, "pairs": len(pairs), "old_control_overlap": overlap,
                      "old_control_overlap_pct": overlap / 1410 * 100})
    candidates[order] = pairs

# Save the closest deterministic implementation, but retain all overlap results.
best = max(summaries, key=lambda x: x["old_control_overlap"])["treated_order"]
pairs = candidates[best]
e["new_weight"] = 0
e["pair_id"] = np.nan
e["matched_role"] = ""
for pid, ti, ci, dist in pairs:
    e.loc[[ti, ci], "new_weight"] = 1
    e.loc[[ti, ci], "pair_id"] = pid
    e.loc[ti, "matched_role"] = "treated"
    e.loc[ci, "matched_role"] = "control"

out = d.merge(e[["name_key", "pscore_reestimated", "common_support", "new_weight", "pair_id", "matched_role"]], on="name_key", how="left")
out = out.merge(old[["name_key", "old_weight", "old_vd"]], on="name_key", how="left")
out["old_weight"] = out.old_weight.fillna(0).astype("int8")
out["same_membership"] = (out.old_weight.eq(out.new_weight.fillna(0))).astype("int8")

pd.DataFrame(summaries).to_csv(ROOT + r"\psm_order_overlap_summary.csv", index=False, encoding="utf-8-sig")
out.to_csv(ROOT + r"\psm_reestimated_membership_comparison.csv", index=False, encoding="utf-8-sig")
stata_out = out.rename(columns={"Organization Name": "Organization_Name", "Organization Name URL": "Organization_Name_URL"}).copy()
for col in stata_out.select_dtypes(include="object"):
    stata_out[col] = stata_out[col].where(stata_out[col].notna(), "").astype(str)
stata_out.to_stata(ROOT + r"\psm_reestimated_full7316.dta", write_index=False, version=118)
stata_out.loc[stata_out.new_weight.fillna(0).eq(1)].to_stata(
    ROOT + r"\psm_reestimated_matched_2820.dta", write_index=False, version=118
)
stata_out.loc[stata_out.old_weight.eq(1)].to_stata(
    ROOT + r"\historical_weight_matched_2820.dta", write_index=False, version=118
)

print(pd.DataFrame(summaries).to_string(index=False))
print("best_order", best)
print("historical_weight", int(old.old_weight.sum()), "historical treated", len(old_treated), "historical controls", len(old_controls))
print("eligible", len(e), "treated", int(e.VD_Backing.sum()), "controls", int((e.VD_Backing == 0).sum()))
print("common support", int(e.common_support.sum()), "treated", int(((e.VD_Backing == 1) & (e.common_support == 1)).sum()))
print("new matched", int(e.new_weight.sum()), "treated", int(((e.new_weight == 1) & (e.VD_Backing == 1)).sum()))
print("score range", float(e.pscore_reestimated.min()), float(e.pscore_reestimated.max()))
hist = out[out.old_weight.eq(1)]
print("historical mapped", len(hist), "treated by source", int(hist.VD_Backing.sum()))
print("same full membership", int(out.same_membership.sum()), "of", len(out))
