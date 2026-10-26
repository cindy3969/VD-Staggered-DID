from pathlib import Path
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
import pandas as pd,re
ROOT=Path('C:/Users/user/Documents/VD/outputs/vd_intensity')
source=Path(r'C:/Users/user/iCloudDrive/5학기/연구미팅/VD PPT/VD Table 20260928.docx')
res=pd.read_csv(ROOT/'tables_3_1_to_3_7/regression_results_long.csv')
stats=pd.read_csv(ROOT/'tables_3_1_to_3_7/model_statistics.csv')
doc=Document(source)
terms=['dose_post','dose_post_prior','dose_post_early','post_prior','post_early','post','prior_patent_l1','lag_early','ln1p_cum_funding_l1','ln1p_rounds_l1']
def setcell(cell,text):
 p=cell.paragraphs[0]
 runs=p.runs
 if runs:
  runs[0].text=text
  for r in runs[1:]:r.text=''
 else:p.add_run(text)
 for p2 in cell.paragraphs[1:]:
  for r in p2.runs:r.text=''
count=0;active=None;checks=[]
for el in doc.element.body:
 if el.tag==qn('w:p'):
  m=re.match(r'Table 3-([1-8])\.',Paragraph(el,doc).text)
  if m:active=int(m[1])
 elif el.tag==qn('w:tbl') and active:
  t=Table(el,doc);n=active;assert len(t.rows)==27
  for j,term in enumerate(terms):
   for model in range(1,5):
    z=res[(res.table==n)&(res.model==model)&(res.term==term)]
    btxt=ptxt=''
    if len(z):
     assert len(z)==1
     r=z.iloc[0];stars='***' if r.p<.01 else '**' if r.p<.05 else '*' if r.p<.1 else ''
     btxt=f'{r.b:.3f}'+stars;ptxt=f'({r.p:.3f})'
    setcell(t.cell(1+2*j,model),btxt);setcell(t.cell(2+2*j,model),ptxt)
    checks.extend([(n,1+2*j,model,btxt),(n,2+2*j,model,ptxt)])
  for model in range(1,5):
   s=stats[(stats.table==n)&(stats.model==model)].iloc[0]
   for row,val in [(24,f'{int(s.N):,}'),(25,f'{int(s.firms):,}'),(26,f'{s.within_r2:.3f}')]:
    setcell(t.cell(row,model),val);checks.append((n,row,model,val))
  count+=1;active=None
assert count==8
out=ROOT/'VD Table_20261019.docx';doc.save(out)
# Independent reopening and comparison of every updated numeric cell.
reopened=Document(out);tables={};active=None
for el in reopened.element.body:
 if el.tag==qn('w:p'):
  m=re.match(r'Table 3-([1-8])\.',Paragraph(el,reopened).text)
  if m:active=int(m[1])
 elif el.tag==qn('w:tbl') and active:
  tables[active]=Table(el,reopened);active=None
for n,row,col,value in checks:assert tables[n].cell(row,col).text==value
print('Updated tables:',count,'verified numeric cells:',len(checks));print(out)
