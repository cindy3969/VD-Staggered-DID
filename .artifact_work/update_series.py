import json
from collections import Counter, defaultdict
from copy import copy
from openpyxl import load_workbook

INPUT = r"C:\Users\user\iCloudDrive\5학기\연구미팅\20240108-2010 to 2015 firms\2010 2015 Crunchbase deals merged_v4.xlsx"
OUTPUT = r"C:\Users\user\Documents\VD\2010 2015 Crunchbase deals merged_v4_series-updated.xlsx"
SUMMARY = r"C:\Users\user\Documents\VD\.artifact_work\series_update_summary.json"
UNKNOWN = "Venture - Series Unknown"
LABELS = ["Pre-Seed", "Seed"] + [f"Series {chr(65+i)}" for i in range(26)]
TO_ORD = {label: i - 1 for i, label in enumerate(LABELS)}

def label_for(ord_):
    if ord_ <= -1:
        return "Pre-Seed"
    if ord_ == 0:
        return "Seed"
    if ord_ <= 26:
        return f"Series {chr(64+ord_)}"
    return "Series Z+"

wb = load_workbook(INPUT, data_only=False)
ws = wb.active
groups = defaultdict(list)
unknown_rows = []

for r in range(3, ws.max_row + 1):
    org_url = ws.cell(r, 2).value
    org_name = ws.cell(r, 4).value
    found_year = ws.cell(r, 7).value
    key = ("url", org_url) if org_url else ("name_year", org_name, found_year)
    rec = {
        "row": r,
        "deal": ws.cell(r, 5).value,
        "date": ws.cell(r, 8).value,
        "type": ws.cell(r, 12).value,
    }
    groups[key].append(rec)
    if rec["type"] == UNKNOWN:
        unknown_rows.append(r)

assignments = {}
methods = Counter()
for key, rows in groups.items():
    rows.sort(key=lambda x: (x["date"] is None, x["date"], x["deal"] or 10**9, x["row"]))
    formal = [x for x in rows if x["type"] in TO_ORD or x["type"] == UNKNOWN]
    pos = 0
    while pos < len(formal):
        if formal[pos]["type"] != UNKNOWN:
            pos += 1
            continue
        start = pos
        while pos < len(formal) and formal[pos]["type"] == UNKNOWN:
            pos += 1
        end = pos - 1
        run = formal[start:pos]
        prev_ord = TO_ORD.get(formal[start - 1]["type"]) if start > 0 else None
        next_ord = TO_ORD.get(formal[pos]["type"]) if pos < len(formal) else None
        m = len(run)
        inferred = []
        if prev_ord is not None and next_ord is not None:
            if next_ord >= prev_ord:
                inferred = [min(prev_ord + j, next_ord) for j in range(1, m + 1)]
                method = "lead-lag bounded sequence"
            else:
                inferred = [prev_ord] * m
                method = "lead series retained due to reverse lag"
        elif prev_ord is not None:
            inferred = [prev_ord + j for j in range(1, m + 1)]
            method = "advanced from lead series"
        elif next_ord is not None:
            raw = [next_ord - (m - j + 1) for j in range(1, m + 1)]
            if min(raw) >= 0:
                inferred = raw
                method = "backcast from lag series"
            else:
                inferred = list(range(0, m))
                method = "round order because lag backcast precedes Seed"
        else:
            inferred = list(range(0, m))
            method = "formal venture-round order"
        for rec, ord_ in zip(run, inferred):
            assignments[rec["row"]] = (label_for(ord_), method)
            methods[method] += 1

if len(assignments) != len(unknown_rows):
    missing = sorted(set(unknown_rows) - set(assignments))
    raise RuntimeError(f"Assignment mismatch: unknown={len(unknown_rows)}, assigned={len(assignments)}, missing={missing[:20]}")

style_ids_before = {r: ws.cell(r, 12).style_id for r in unknown_rows}
for r, (label, method) in assignments.items():
    ws.cell(r, 12).value = label

wb.save(OUTPUT)

check = load_workbook(OUTPUT, read_only=True, data_only=False)
cws = check.active
remaining = 0
distribution = Counter()
style_mismatch = 0
target_rows = set(unknown_rows)
for r, row in enumerate(cws.iter_rows(min_row=3), 3):
    if r not in target_rows:
        continue
    cell = row[11]
    val = cell.value
    if val == UNKNOWN:
        remaining += 1
    distribution[val] += 1
    if cell._style_id != style_ids_before[r]:
        style_mismatch += 1

summary = {
    "input": INPUT,
    "output": OUTPUT,
    "sheet": ws.title,
    "updated_column": "L",
    "unknown_before": len(unknown_rows),
    "unknown_after": remaining,
    "inferred_distribution": dict(distribution.most_common()),
    "methods": dict(methods),
    "style_mismatch_count": style_mismatch,
    "row_count": ws.max_row,
    "column_count": ws.max_column,
}
with open(SUMMARY, "w", encoding="utf-8") as f:
    json.dump(summary, f, ensure_ascii=False, indent=2, default=str)
print(json.dumps(summary, ensure_ascii=False, indent=2, default=str))
