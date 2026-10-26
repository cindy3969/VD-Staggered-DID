import json
import re
import zipfile
from collections import Counter, defaultdict
from openpyxl import load_workbook

INPUT = r"C:\Users\user\iCloudDrive\5학기\연구미팅\20240108-2010 to 2015 firms\2010 2015 Crunchbase deals merged_v4.xlsx"
OUTPUT = r"C:\Users\user\Documents\VD\2010 2015 Crunchbase deals merged_v4_series-updated.xlsx"
SUMMARY = r"C:\Users\user\Documents\VD\.artifact_work\series_update_summary.json"
SHEET_XML = "xl/worksheets/sheet1.xml"
UNKNOWN = "Venture - Series Unknown"
LABELS = ["Pre-Seed", "Seed"] + [f"Series {chr(65+i)}" for i in range(26)]
TO_ORD = {label: i - 1 for i, label in enumerate(LABELS)}

def label_for(ord_):
    if ord_ <= -1:
        return "Pre-Seed"
    if ord_ == 0:
        return "Seed"
    if ord_ <= 26:
        return f"Series {chr(64+ord_)}"
    return "Series Z+"

wb = load_workbook(INPUT, read_only=True, data_only=False)
ws = wb.active
groups = defaultdict(list)
unknown_rows = []
for r, vals in enumerate(ws.iter_rows(min_row=3, values_only=True), 3):
    org_url, org_name, found_year = vals[1], vals[3], vals[6]
    key = ("url", org_url) if org_url else ("name_year", org_name, found_year)
    rec = {"row": r, "deal": vals[4], "date": vals[7], "type": vals[11]}
    groups[key].append(rec)
    if rec["type"] == UNKNOWN:
        unknown_rows.append(r)

assignments = {}
methods = Counter()
for key, rows in groups.items():
    rows.sort(key=lambda x: (x["date"] is None, x["date"], x["deal"] or 10**9, x["row"]))
    formal = [x for x in rows if x["type"] in TO_ORD or x["type"] == UNKNOWN]
    pos = 0
    while pos < len(formal):
        if formal[pos]["type"] != UNKNOWN:
            pos += 1
            continue
        start = pos
        while pos < len(formal) and formal[pos]["type"] == UNKNOWN:
            pos += 1
        run = formal[start:pos]
        prev_ord = TO_ORD.get(formal[start - 1]["type"]) if start > 0 else None
        next_ord = TO_ORD.get(formal[pos]["type"]) if pos < len(formal) else None
        m = len(run)
        if prev_ord is not None and next_ord is not None:
            if next_ord >= prev_ord:
                inferred = [min(prev_ord + j, next_ord) for j in range(1, m + 1)]
                method = "lead-lag bounded sequence"
            else:
                inferred = [prev_ord] * m
                method = "lead series retained due to reverse lag"
        elif prev_ord is not None:
            inferred = [prev_ord + j for j in range(1, m + 1)]
            method = "advanced from lead series"
        elif next_ord is not None:
            raw = [next_ord - (m - j + 1) for j in range(1, m + 1)]
            if min(raw) >= 0:
                inferred = raw
                method = "backcast from lag series"
            else:
                inferred = list(range(0, m))
                method = "round order because lag backcast precedes Seed"
        else:
            inferred = list(range(0, m))
            method = "formal venture-round order"
        for rec, ord_ in zip(run, inferred):
            assignments[rec["row"]] = (label_for(ord_), method)
            methods[method] += 1

if len(assignments) != len(unknown_rows):
    raise RuntimeError(f"Assignment mismatch: {len(unknown_rows)} unknown versus {len(assignments)} assigned")

cell_pattern = re.compile(rb'<c r="L(\d+)"([^>]*)>(.*?)</c>')
changed = 0

def replace_cell(match):
    global changed
    row = int(match.group(1))
    if row not in assignments:
        return match.group(0)
    label = assignments[row][0]
    attrs = re.sub(rb'\s+t="[^"]*"', b'', match.group(2))
    text = label.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    changed += 1
    return f'<c r="L{row}"'.encode() + attrs + b' t="inlineStr"><is><t>' + text.encode("utf-8") + b'</t></is></c>'

with zipfile.ZipFile(INPUT, "r") as zin, zipfile.ZipFile(OUTPUT, "w") as zout:
    for info in zin.infolist():
        data = zin.read(info.filename)
        if info.filename == SHEET_XML:
            data = cell_pattern.sub(replace_cell, data)
        zout.writestr(info, data)

if changed != len(assignments):
    raise RuntimeError(f"XML replacement mismatch: expected {len(assignments)}, changed {changed}")

distribution = Counter(label for label, _ in assignments.values())
summary = {
    "input": INPUT,
    "output": OUTPUT,
    "sheet": ws.title,
    "updated_column": "L",
    "unknown_before": len(unknown_rows),
    "xml_cells_changed": changed,
    "inferred_distribution": dict(distribution.most_common()),
    "methods": dict(methods),
    "row_count": ws.max_row,
    "column_count": ws.max_column,
}
with open(SUMMARY, "w", encoding="utf-8") as f:
    json.dump(summary, f, ensure_ascii=False, indent=2)
print(json.dumps(summary, ensure_ascii=False, indent=2))
