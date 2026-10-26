from copy import copy
from openpyxl import Workbook, load_workbook

SOURCE = r"C:\Users\user\Documents\VD\2010 2015 Crunchbase deals merged_v4_series-updated.xlsx"
PREVIEW = r"C:\Users\user\Documents\VD\.artifact_work\series_preview.xlsx"
src = load_workbook(SOURCE, read_only=False, data_only=False)
s = src.active
wb = Workbook()
w = wb.active
w.title = s.title
for row in s.iter_rows(min_row=1, max_row=35, min_col=1, max_col=14):
    for c in row:
        d = w.cell(c.row, c.column, c.value)
        if c.has_style:
            d.font = copy(c.font)
            d.fill = copy(c.fill)
            d.border = copy(c.border)
            d.alignment = copy(c.alignment)
            d.protection = copy(c.protection)
        if c.number_format:
            d.number_format = c.number_format
for col in range(1, 15):
    letter = w.cell(1, col).column_letter
    w.column_dimensions[letter].width = s.column_dimensions[letter].width
for r in range(1, 36):
    w.row_dimensions[r].height = s.row_dimensions[r].height
w.freeze_panes = s.freeze_panes
w.sheet_view.showGridLines = s.sheet_view.showGridLines
wb.save(PREVIEW)
print(PREVIEW)
