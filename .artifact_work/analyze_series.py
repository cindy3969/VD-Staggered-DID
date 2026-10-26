from collections import Counter, defaultdict
from openpyxl import load_workbook

PATH = r"C:\Users\user\iCloudDrive\5학기\연구미팅\20240108-2010 to 2015 firms\2010 2015 Crunchbase deals merged_v4.xlsx"
UNKNOWN = "Venture - Series Unknown"
LABELS = ["Pre-Seed", "Seed"] + [f"Series {chr(65+i)}" for i in range(26)]
TO_ORD = {label: i - 1 for i, label in enumerate(LABELS)}

def label_for(ord_):
    if ord_ <= -1:
        return "Pre-Seed"
    if ord_ == 0:
        return "Seed"
    if ord_ <= 26:
        return f"Series {chr(64+ord_)}"
    return "Series Z+"

wb = load_workbook(PATH, read_only=True, data_only=False)
ws = wb.active
groups = defaultdict(list)
for r, vals in enumerate(ws.iter_rows(min_row=3, values_only=True), 3):
    org = vals[1]
    if not org:
        continue
    groups[org].append({"row": r, "deal": vals[4], "date": vals[7], "type": vals[11]})

result = {}
methods = Counter()
conflicts = []
for org, rows in groups.items():
    rows.sort(key=lambda x: (x["date"] is None, x["date"], x["deal"] or 10**9, x["row"]))
    formal = [x for x in rows if x["type"] in TO_ORD or x["type"] == UNKNOWN]
    for pos, x in enumerate(formal):
        if x["type"] != UNKNOWN:
            continue
        prev_i = next((j for j in range(pos - 1, -1, -1) if formal[j]["type"] in TO_ORD), None)
        next_i = next((j for j in range(pos + 1, len(formal)) if formal[j]["type"] in TO_ORD), None)
        by_order = pos - (1 if formal and formal[0]["type"] == "Pre-Seed" else 0)
        by_order = max(0, by_order)
        inferred = None
        method = None
        if prev_i is not None and next_i is not None:
            p = TO_ORD[formal[prev_i]["type"]]
            n = TO_ORD[formal[next_i]["type"]]
            from_prev = p + (pos - prev_i)
            from_next = n - (next_i - pos)
            if from_prev == from_next:
                inferred = from_prev
                method = "interpolated from consistent lead and lag series"
            else:
                inferred = from_prev
                method = "lead series; lag conflict noted"
                conflicts.append((x["row"], org, formal[prev_i]["type"], formal[next_i]["type"], label_for(from_prev), label_for(from_next)))
        elif prev_i is not None:
            p = TO_ORD[formal[prev_i]["type"]]
            inferred = p + (pos - prev_i)
            method = "advanced from lead series"
        elif next_i is not None:
            n = TO_ORD[formal[next_i]["type"]]
            inferred = n - (next_i - pos)
            if inferred < 0:
                inferred = by_order
                method = "round order; lag series would precede Seed"
            else:
                method = "backcast from lag series"
        else:
            inferred = by_order
            method = "round order"
        result[x["row"]] = (label_for(inferred), method)
        methods[method] += 1

print("UNKNOWN", len(result))
print("INFERRED", Counter(v[0] for v in result.values()).most_common())
print("METHODS", methods)
print("CONFLICTS", len(conflicts))
for item in conflicts[:50]:
    print("CONFLICT", item)
