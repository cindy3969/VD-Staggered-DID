import json
from itertools import zip_longest
from openpyxl import load_workbook

SOURCE = r"C:\Users\user\iCloudDrive\5학기\연구미팅\20240108-2010 to 2015 firms\2010 2015 Crunchbase deals merged_v4.xlsx"
OUTPUT = r"C:\Users\user\Documents\VD\2010 2015 Crunchbase deals merged_v4_series-updated.xlsx"
UNKNOWN = "Venture - Series Unknown"

src = load_workbook(SOURCE, read_only=True, data_only=False)
out = load_workbook(OUTPUT, read_only=True, data_only=False)
assert src.sheetnames == out.sheetnames
s1, s2 = src.active, out.active
diff_count = 0
bad_diff_count = 0
formula_src = 0
formula_out = 0
unknown_out = 0
error_values = []
diff_samples = []

for r_idx, (row1, row2) in enumerate(zip_longest(s1.iter_rows(), s2.iter_rows()), 1):
    if row1 is None or row2 is None:
        raise RuntimeError("Row count mismatch")
    for c_idx, (c1, c2) in enumerate(zip_longest(row1, row2), 1):
        if c1 is None or c2 is None:
            raise RuntimeError(f"Column count mismatch at row {r_idx}")
        v1, v2 = c1.value, c2.value
        formula_src += isinstance(v1, str) and v1.startswith("=")
        formula_out += isinstance(v2, str) and v2.startswith("=")
        if c_idx == 12 and v2 == UNKNOWN:
            unknown_out += 1
        if isinstance(v2, str) and v2 in {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}:
            error_values.append((r_idx, c_idx, v2))
        if v1 != v2:
            diff_count += 1
            valid = c_idx == 12 and v1 == UNKNOWN and isinstance(v2, str) and (v2 in {"Pre-Seed", "Seed"} or v2.startswith("Series "))
            if not valid:
                bad_diff_count += 1
            if len(diff_samples) < 20:
                diff_samples.append((r_idx, c_idx, v1, v2, valid))

result = {
    "source_dimensions": [s1.max_row, s1.max_column],
    "output_dimensions": [s2.max_row, s2.max_column],
    "sheet_names_equal": src.sheetnames == out.sheetnames,
    "difference_count": diff_count,
    "unexpected_difference_count": bad_diff_count,
    "formula_count_source": formula_src,
    "formula_count_output": formula_out,
    "unknown_remaining_output": unknown_out,
    "formula_error_values": error_values[:50],
    "difference_samples": diff_samples,
}
print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
